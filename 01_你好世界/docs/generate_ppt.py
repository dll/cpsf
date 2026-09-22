"""
第1阶段：表达式 —— C语言的灵魂 - PPT课件生成脚本
基于 python-pptx 生成图文并茂的教学课件
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ==========================================
# 配色方案（亮色版）
# ==========================================
COLORS = {
    'primary': RGBColor(0x2C, 0x5F, 0x8B),    # 主色：深蓝
    'secondary': RGBColor(0x2E, 0xCC, 0xFA),   # 辅助色：青
    'accent': RGBColor(0x9C, 0x27, 0xB0),      # 强调色：紫
    'highlight': RGBColor(0xFF, 0xA0, 0x00),   # 重点色：金
    'success': RGBColor(0x4C, 0xAF, 0x50),     # 成功绿
    'danger': RGBColor(0xEF, 0x53, 0x50),      # 危险红
    'bg_light': RGBColor(0xF5, 0xF7, 0xFA),    # 浅色背景
    'bg_dark': RGBColor(0x1A, 0x23, 0x32),     # 深色背景
    'text_dark': RGBColor(0x33, 0x33, 0x33),   # 深色文字
    'text_light': RGBColor(0xFF, 0xFF, 0xFF),  # 浅色文字
    'text_gray': RGBColor(0x66, 0x66, 0x66),   # 灰色文字
}

FONT_TITLE = '微软雅黑'
FONT_BODY = '微软雅黑'
FONT_CODE = 'Consolas'


def set_slide_bg(slide, color):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height, text, font_size=18,
                font_name=FONT_BODY, color=COLORS['text_dark'],
                bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.name = font_name
    run.font.color.rgb = color
    run.font.bold = bold
    return txBox


def add_rounded_rect(slide, left, top, width, height, fill_color, line_color=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
    else:
        shape.line.fill.background()
    return shape


def add_card(slide, left, top, width, height, title, content,
             title_color=COLORS['primary']):
    card = add_rounded_rect(slide, left, top, width, height,
                            RGBColor(0xFF, 0xFF, 0xFF),
                            RGBColor(0xE0, 0xE0, 0xE0))
    # 顶部色条
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, left, top, width, Emu(50800)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = title_color
    bar.line.fill.background()
    # 标题
    add_textbox(slide, left + Inches(0.3), top + Inches(0.2),
                width - Inches(0.6), Inches(0.5),
                title, font_size=16, bold=True, color=title_color)
    # 内容
    add_textbox(slide, left + Inches(0.3), top + Inches(0.7),
                width - Inches(0.6), height - Inches(0.9),
                content, font_size=13, color=COLORS['text_dark'])


def add_code_block(slide, left, top, width, height, code_text):
    bg = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(0x1E, 0x1E, 0x1E)
    bg.line.fill.background()
    txBox = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15),
                                     width - Inches(0.4), height - Inches(0.3))
    tf = txBox.text_frame
    tf.word_wrap = True
    lines = code_text.strip().split('\n')
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = line
        run.font.size = Pt(13)
        run.font.name = FONT_CODE
        run.font.color.rgb = RGBColor(0xD4, 0xD4, 0xD4)


def add_table(slide, left, top, width, height, data, header_color=COLORS['primary']):
    rows = len(data)
    cols = len(data[0])
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table
    for i in range(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.text = str(data[i][j])
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = FONT_BODY
                    run.font.size = Pt(13)
                    if i == 0:
                        run.font.bold = True
                        run.font.color.rgb = COLORS['text_light']
                    else:
                        run.font.color.rgb = COLORS['text_dark']
            if i == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = header_color
            elif i % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0xF8, 0xF9, 0xFA)
    return table


def add_uml_box(slide, left, top, width, height, text, fill_color, text_color=COLORS['text_light']):
    """添加UML风格的方框"""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = RGBColor(0x33, 0x33, 0x33)
    shape.line.width = Pt(1.5)
    # 文字
    add_textbox(slide, left, top + height / 2 - Inches(0.3),
                width, Inches(0.6),
                text, font_size=14, bold=True,
                color=text_color, align=PP_ALIGN.CENTER)
    return shape


def add_arrow_down(slide, left, top, height, color=COLORS['secondary']):
    """添加向下箭头"""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.DOWN_ARROW, left, top, Inches(0.4), height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def create_ppt():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # ==========================================
    # 第1页：封面
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, COLORS['bg_dark'])

    # 装饰：UML类图线稿感的装饰元素
    for i in range(5):
        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.5 + i * 2.5), Inches(1),
            Inches(0.5), Inches(0.5)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = RGBColor(0x2A, 0x3A, 0x50)
        line.line.color.rgb = RGBColor(0x3D, 0x5A, 0x80)
        line.line.width = Pt(1)

    # 竖线装饰
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1), Inches(2.8), Inches(0.08), Inches(2)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = COLORS['highlight']
    line.line.fill.background()

    # 主标题
    add_textbox(slide, Inches(1.5), Inches(2.5), Inches(10), Inches(1.2),
                '第1讲：表达式', font_size=54, bold=True,
                color=COLORS['text_light'])

    # 副标题
    add_textbox(slide, Inches(1.5), Inches(3.7), Inches(10), Inches(0.8),
                '—— C语言的灵魂', font_size=28,
                color=COLORS['secondary'])

    # 系列名称
    add_textbox(slide, Inches(1.5), Inches(5), Inches(10), Inches(0.6),
                '从程序员到软件框架 —— C语言插件框架演进之旅', font_size=16,
                color=COLORS['text_gray'])

    # 阶段标识
    add_textbox(slide, Inches(10.5), Inches(6.2), Inches(2), Inches(0.5),
                '01 / 12', font_size=20, bold=True,
                color=COLORS['highlight'], align=PP_ALIGN.RIGHT)

    # ==========================================
    # 第2页：导入 - C语言是怎么来的
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, COLORS['bg_light'])

    add_rounded_rect(slide, Inches(0.5), Inches(0.4), Inches(12.3), Inches(0.8),
                     COLORS['primary'])
    add_textbox(slide, Inches(0.8), Inches(0.5), Inches(12), Inches(0.6),
                '📜 导入：C语言是怎么来的？', font_size=24, bold=True,
                color=COLORS['text_light'])

    # 时间线
    timeline_items = [
        ('1960s', 'BCPL\n语言', COLORS['text_gray']),
        ('1970', 'B语言\nKen Thompson', COLORS['secondary']),
        ('1972', 'C语言诞生\nDennis Ritchie', COLORS['highlight']),
        ('1973', 'UNIX用C\n重写', COLORS['success']),
        ('1978', 'K&R C\n经典教材', COLORS['accent']),
    ]

    # 时间轴线
    timeline = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1), Inches(2.8), Inches(11.3), Inches(0.08)
    )
    timeline.fill.solid()
    timeline.fill.fore_color.rgb = COLORS['primary']
    timeline.line.fill.background()

    for i, (year, desc, color) in enumerate(timeline_items):
        left = Inches(1 + i * 2.6)
        # 节点圆点
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, left + Inches(0.8), Inches(2.6),
            Inches(0.5), Inches(0.5)
        )
        dot.fill.solid()
        dot.fill.fore_color.rgb = color
        dot.line.color.rgb = COLORS['text_light']
        dot.line.width = Pt(2)
        # 年份
        add_textbox(slide, left, Inches(1.8), Inches(2.1), Inches(0.5),
                    year, font_size=18, bold=True,
                    color=COLORS['primary'], align=PP_ALIGN.CENTER)
        # 描述
        add_textbox(slide, left, Inches(3.2), Inches(2.1), Inches(1),
                    desc, font_size=12,
                    color=COLORS['text_dark'], align=PP_ALIGN.CENTER)

    # 底部故事
    add_rounded_rect(slide, Inches(1), Inches(4.8), Inches(11.3), Inches(2),
                     RGBColor(0xFF, 0xF8, 0xE1), RGBColor(0xFF, 0xA0, 0x00))
    add_textbox(slide, Inches(1.3), Inches(5), Inches(10.7), Inches(0.5),
                '🎯 一个"副产品"改变了世界', font_size=18, bold=True, color=COLORS['highlight'])
    add_textbox(slide, Inches(1.3), Inches(5.5), Inches(10.7), Inches(1.2),
                '1972年，贝尔实验室，Dennis Ritchie 为了写 UNIX 操作系统，在B语言基础上发明了C语言。\n'
                'C语言为什么叫C？因为它的前身是B语言（BCPL的首字母），C是B的下一个字母。\n'
                '程序员起名字，就是这么朴实无华。',
                font_size=14, color=COLORS['text_dark'])

    # ==========================================
    # 第3页：C语言为什么伟大
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, COLORS['bg_light'])

    add_rounded_rect(slide, Inches(0.5), Inches(0.4), Inches(12.3), Inches(0.8),
                     COLORS['accent'])
    add_textbox(slide, Inches(0.8), Inches(0.5), Inches(12), Inches(0.6),
                '🏆 C语言为什么伟大？', font_size=24, bold=True,
                color=COLORS['text_light'])

    # 三种语言对比
    langs = [
        ('汇编语言', '最快、最直接\n完全操控硬件', '难写、难读\n不可移植', COLORS['danger']),
        ('C语言', '接近机器效率\n又有人类可读', '手动管理内存\n容易出bug', COLORS['highlight']),
        ('高级语言', '好写好读\n功能强大', '运行效率低\n离硬件远', COLORS['success']),
    ]

    for i, (name, pros, cons, color) in enumerate(langs):
        left = Inches(0.8 + i * 4.1)
        add_card(slide, left, Inches(1.6), Inches(3.8), Inches(2.5),
                 name, '', title_color=color)
        add_textbox(slide, left + Inches(0.3), Inches(2.3), Inches(3.2), Inches(0.4),
                    '✅ 优点', font_size=13, bold=True, color=COLORS['success'])
        add_textbox(slide, left + Inches(0.3), Inches(2.7), Inches(3.2), Inches(0.7),
                    pros, font_size=12, color=COLORS['text_dark'])
        add_textbox(slide, left + Inches(0.3), Inches(3.4), Inches(3.2), Inches(0.4),
                    '❌ 缺点', font_size=13, bold=True, color=COLORS['danger'])
        add_textbox(slide, left + Inches(0.3), Inches(3.8), Inches(3.2), Inches(0.5),
                    cons, font_size=12, color=COLORS['text_dark'])

    # 底部：C语言的江湖地位
    add_rounded_rect(slide, Inches(0.8), Inches(4.5), Inches(11.7), Inches(2.5),
                     RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xE0, 0xE0, 0xE0))
    add_textbox(slide, Inches(1.1), Inches(4.7), Inches(11), Inches(0.5),
                '💎 C语言的江湖地位', font_size=18, bold=True, color=COLORS['primary'])

    achievements = [
        '🏆 UNIX/Linux 内核',
        '🏆 Windows 内核（主要）',
        '🏆 MySQL / SQLite / Redis',
        '🏆 Python / PHP / Ruby 解释器',
        '🏆 几乎所有嵌入式系统',
    ]
    for i, item in enumerate(achievements):
        col = i % 3
        row = i // 3
        add_textbox(slide, Inches(1.1 + col * 3.8), Inches(5.3 + row * 0.7),
                    Inches(3.5), Inches(0.5),
                    item, font_size=13, color=COLORS['text_dark'])

    # ==========================================
    # 第4页：表达式 —— C语言的灵魂
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, COLORS['bg_light'])

    add_rounded_rect(slide, Inches(0.5), Inches(0.4), Inches(12.3), Inches(0.8),
                     COLORS['secondary'])
    add_textbox(slide, Inches(0.8), Inches(0.5), Inches(12), Inches(0.6),
                '🧩 表达式 —— C语言的灵魂', font_size=24, bold=True,
                color=COLORS['text_light'])

    # 核心定义
    add_rounded_rect(slide, Inches(1), Inches(1.5), Inches(11.3), Inches(1),
                     COLORS['primary'])
    add_textbox(slide, Inches(1.3), Inches(1.7), Inches(10.7), Inches(0.6),
                '表达式（Expression）= 能算出值的代码片段',
                font_size=22, bold=True, color=COLORS['text_light'], align=PP_ALIGN.CENTER)

    # 主角代码
    add_code_block(slide, Inches(3.5), Inches(2.8), Inches(6.3), Inches(1.2),
                   'printf("Hello, Plugin Framework!\\n");')

    # 表达式类型
    types = [
        ('常量表达式', '42、"Hello"', COLORS['primary']),
        ('函数调用表达式', 'printf("Hi")', COLORS['secondary']),
        ('算术表达式', '1 + 2 * 3', COLORS['accent']),
        ('赋值表达式', 'a = 5', COLORS['highlight']),
        ('逗号表达式', 'a=1, b=2', COLORS['success']),
    ]

    for i, (name, example, color) in enumerate(types):
        left = Inches(0.5 + i * 2.5)
        add_rounded_rect(slide, left, Inches(4.3), Inches(2.3), Inches(1.2),
                         color)
        add_textbox(slide, left, Inches(4.4), Inches(2.3), Inches(0.4),
                    name, font_size=13, bold=True,
                    color=COLORS['text_light'], align=PP_ALIGN.CENTER)
        add_textbox(slide, left + Inches(0.1), Inches(4.9), Inches(2.1), Inches(0.5),
                    example, font_size=12,
                    color=COLORS['text_light'], align=PP_ALIGN.CENTER,
                    font_name=FONT_CODE)

    # 底部幽默
    add_rounded_rect(slide, Inches(1), Inches(5.8), Inches(11.3), Inches(1.2),
                     RGBColor(0xF3, 0xE5, 0xF5), RGBColor(0x9C, 0x27, 0xB0))
    add_textbox(slide, Inches(1.3), Inches(6), Inches(10.7), Inches(0.4),
                '🎭 冷知识', font_size=15, bold=True, color=COLORS['accent'])
    add_textbox(slide, Inches(1.3), Inches(6.4), Inches(10.7), Inches(0.5),
                'C语言的表达式有多疯狂？连 a = b = c = 5 都是合法的——因为赋值也是表达式，有返回值！'
                '别的语言："赋值是语句。" C语言："不，我就是要玩出花来。"',
                font_size=13, color=COLORS['text_dark'])

    # ==========================================
    # 第5页：程序层次图（UML风格）
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, COLORS['bg_light'])

    add_rounded_rect(slide, Inches(0.5), Inches(0.4), Inches(12.3), Inches(0.8),
                     COLORS['primary'])
    add_textbox(slide, Inches(0.8), Inches(0.5), Inches(12), Inches(0.6),
                '🏗️ 从表达式到程序：层次结构图', font_size=24, bold=True,
                color=COLORS['text_light'])

    # UML风格的层次结构（从上到下：大→小）
    levels = [
        ('程序 Program', 'hello.exe', COLORS['primary'], Inches(4.5), Inches(1.3)),
        ('函数 Function', 'main()、printf()', COLORS['secondary'], Inches(4), Inches(1.1)),
        ('语句 Statement', 'printf("Hi");', COLORS['accent'], Inches(3.5), Inches(0.9)),
        ('表达式 Expression', '1+2、printf("Hi")', COLORS['highlight'], Inches(3), Inches(0.8)),
        ('常量/变量/运算符', '42、a、+', COLORS['success'], Inches(2.5), Inches(0.7)),
    ]

    center_x = Inches(6.666)

    for i, (name, example, color, width, height) in enumerate(levels):
        left = center_x - width / 2
        top = Inches(1.4 + i * 1.15)
        add_uml_box(slide, left, top, width, height, f'{name}\n（{example}）', color)
        if i < len(levels) - 1:
            # 向下箭头
            arrow_left = center_x - Inches(0.2)
            arrow_top = top + height
            add_arrow_down(slide, arrow_left, arrow_top, Inches(0.2), color=COLORS['text_gray'])

    # 右侧说明
    add_card(slide, Inches(9.5), Inches(1.5), Inches(3.3), Inches(4.5),
             '💡 关键洞察',
             '表达式是最底层的砖块。\n\n'
             '砖块虽小，\n但组合方式对了，\n就能盖出摩天大楼。\n\n'
             '我们这个系列，\n就是看这些"砖块"\n怎么一步步变成\n"框架"这座大厦的。',
             title_color=COLORS['highlight'])

    # ==========================================
    # 第6页：第一行代码详解
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, COLORS['bg_light'])

    add_rounded_rect(slide, Inches(0.5), Inches(0.4), Inches(12.3), Inches(0.8),
                     COLORS['accent'])
    add_textbox(slide, Inches(0.8), Inches(0.5), Inches(12), Inches(0.6),
                '📝 我们的第一行代码', font_size=24, bold=True,
                color=COLORS['text_light'])

    # 左侧代码
    code = '''#include <stdio.h>

int main(void)
{
    printf("Hello, Plugin Framework!\\n");
    return 0;
}'''
    add_code_block(slide, Inches(0.8), Inches(1.6), Inches(6), Inches(4.5), code)

    # 右侧注释
    annotations = [
        ('① 头文件引入', '#include <stdio.h>\n把标准IO头文件"抄"进来', '📋', COLORS['primary']),
        ('② 主函数入口', 'int main(void)\n程序从这里开始执行', '🚪', COLORS['secondary']),
        ('③ 函数调用表达式', 'printf("...")\n这是一个表达式！有返回值', '🖨️', COLORS['accent']),
        ('④ 返回值表达式', 'return 0;\n返回0 = 一切正常', '✅', COLORS['success']),
    ]

    for i, (title, desc, emoji, color) in enumerate(annotations):
        top = Inches(1.6 + i * 1.15)
        add_card(slide, Inches(7.2), top, Inches(5.3), Inches(1),
                 f'{emoji} {title}', desc, title_color=color)

    # ==========================================
    # 第7页：编译运行流程（UML活动图风格）
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, COLORS['bg_light'])

    add_rounded_rect(slide, Inches(0.5), Inches(0.4), Inches(12.3), Inches(0.8),
                     COLORS['secondary'])
    add_textbox(slide, Inches(0.8), Inches(0.5), Inches(12), Inches(0.6),
                '⚙️ 编译运行流程（UML活动图）', font_size=24, bold=True,
                color=COLORS['text_light'])

    # 流程图
    stages = [
        ('hello.c', '源代码', COLORS['primary']),
        ('预处理', '展开宏和头文件\nhello.i', COLORS['secondary']),
        ('编译', '翻译成汇编\nhello.s', COLORS['accent']),
        ('汇编', '生成目标文件\nhello.o', COLORS['highlight']),
        ('链接', '拼入printf\n+ stdio库', COLORS['success']),
        ('hello.exe', '可执行文件', COLORS['danger']),
    ]

    for i, (title, desc, color) in enumerate(stages):
        left = Inches(0.4 + i * 2.1)
        if i == 0 or i == 5:
            # 起止节点：圆形
            shape = slide.shapes.add_shape(
                MSO_SHAPE.OVAL, left, Inches(2.5), Inches(1.9), Inches(1.5)
            )
        else:
            # 中间节点：圆角矩形
            shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.5), Inches(1.9), Inches(1.5)
            )
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.color.rgb = RGBColor(0x33, 0x33, 0x33)
        shape.line.width = Pt(1.5)

        add_textbox(slide, left, Inches(2.7), Inches(1.9), Inches(0.4),
                    title, font_size=14, bold=True,
                    color=COLORS['text_light'], align=PP_ALIGN.CENTER)
        add_textbox(slide, left + Inches(0.1), Inches(3.1), Inches(1.7), Inches(0.8),
                    desc, font_size=11,
                    color=COLORS['text_light'], align=PP_ALIGN.CENTER)

        if i < 5:
            # 箭头
            arrow_left = left + Inches(1.9)
            arrow_shape = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW, arrow_left, Inches(3.1),
                Inches(0.2), Inches(0.3)
            )
            arrow_shape.fill.solid()
            arrow_shape.fill.fore_color.rgb = COLORS['text_gray']
            arrow_shape.line.fill.background()

    # 重点洞察
    add_rounded_rect(slide, Inches(1), Inches(4.7), Inches(11.3), Inches(2.3),
                     RGBColor(0xFF, 0xF8, 0xE1), RGBColor(0xFF, 0xA0, 0x00))
    add_textbox(slide, Inches(1.3), Inches(4.9), Inches(10.7), Inches(0.5),
                '🎯 重点中的重点', font_size=18, bold=True, color=COLORS['highlight'])
    add_textbox(slide, Inches(1.3), Inches(5.4), Inches(10.7), Inches(1.5),
                '你写的代码里没有 printf 的实现，但程序却能运行它！\n\n'
                '因为「链接器」把标准库里的 printf 和你的代码"拼"在了一起。\n\n'
                '这就是 复用 的起点 —— 用别人写好的代码，做自己的事情。\n'
                '这也是 接口 的起点 —— 你不需要知道实现，只需要知道怎么用。',
                font_size=14, color=COLORS['text_dark'])

    # ==========================================
    # 第8页：AI化辅助
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, COLORS['bg_light'])

    add_rounded_rect(slide, Inches(0.5), Inches(0.4), Inches(12.3), Inches(0.8),
                     COLORS['accent'])
    add_textbox(slide, Inches(0.8), Inches(0.5), Inches(12), Inches(0.6),
                '🤖 AI化辅助：探索表达式的奥秘', font_size=24, bold=True,
                color=COLORS['text_light'])

    # Prompt示例
    add_card(slide, Inches(0.8), Inches(1.5), Inches(11.7), Inches(1.8),
             '💬 Prompt 示例',
             '"我是大三计算机专业学生，刚学完C语言基础。\n'
             '请给我讲讲C语言中「表达式」这个概念的独特之处。\n'
             '为什么说C语言「几乎一切都是表达式」？和Java/Python有什么不同？举3个骚操作例子。"',
             title_color=COLORS['accent'])

    # 三个骚操作
    tricks = [
        ('连锁赋值', 'a = b = c = 5;\n赋值表达式有返回值\n所以可以一路赋过去', '🔗'),
        ('赋值嵌if', 'if ((p=malloc(...))!=NULL)\n边赋值边判断\n一行搞定', '🤝'),
        ('逗号表达式', 'for(i=0,j=10; i<j; i++,j--)\n逗号分隔多个表达式\n取最后一个的值', '🦠'),
    ]

    for i, (title, desc, emoji) in enumerate(tricks):
        left = Inches(0.8 + i * 4.1)
        add_card(slide, left, Inches(3.7), Inches(3.8), Inches(2.5),
                 f'{emoji} {title}', desc, title_color=COLORS['secondary'])

    # 核验提示
    add_rounded_rect(slide, Inches(0.8), Inches(6.5), Inches(11.7), Inches(0.7),
                     RGBColor(0xE8, 0xF5, 0xE9), RGBColor(0x4C, 0xAF, 0x50))
    add_textbox(slide, Inches(1.1), Inches(6.6), Inches(11), Inches(0.5),
                '✅ 核验练习：把三个"骚操作"写成小程序编译运行，验证AI说的对不对',
                font_size=13, bold=True, color=COLORS['success'])

    # ==========================================
    # 第9页：做中学
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, COLORS['bg_light'])

    add_rounded_rect(slide, Inches(0.5), Inches(0.4), Inches(12.3), Inches(0.8),
                     COLORS['highlight'])
    add_textbox(slide, Inches(0.8), Inches(0.5), Inches(12), Inches(0.6),
                '🛠️ 做中学：动手实验', font_size=24, bold=True,
                color=COLORS['text_light'])

    experiments = [
        ('实验1：亲手敲代码', '不要复制粘贴，\n亲手敲出hello.c', '观察：你打错了几次？\n怎么修好的？', '⌨️', COLORS['primary']),
        ('实验2：表达式套娃', '打印printf的返回值\n看看它返回了什么', '思考：返回值是几？\n这说明了什么？', '🎁', COLORS['secondary']),
        ('实验3：故意制造错误', '至少试3种错误写法\n记录编译器报错', '思考：报错信息有用吗？\n你会怎么改进？', '🐛', COLORS['accent']),
    ]

    for i, (title, task, think, emoji, color) in enumerate(experiments):
        left = Inches(0.8 + i * 4.1)
        add_card(slide, left, Inches(1.5), Inches(3.8), Inches(4.5),
                 f'{emoji} {title}', '', title_color=color)
        # 任务
        add_textbox(slide, left + Inches(0.3), Inches(2.3), Inches(3.2), Inches(0.4),
                    '📋 任务：', font_size=14, bold=True, color=color)
        add_textbox(slide, left + Inches(0.3), Inches(2.7), Inches(3.2), Inches(1),
                    task, font_size=13, color=COLORS['text_dark'])
        # 分割线
        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, left + Inches(0.3), Inches(3.9),
            Inches(3.2), Emu(12700)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = RGBColor(0xE0, 0xE0, 0xE0)
        line.line.fill.background()
        # 思考
        add_textbox(slide, left + Inches(0.3), Inches(4.1), Inches(3.2), Inches(0.4),
                    '💭 思考：', font_size=14, bold=True, color=color)
        add_textbox(slide, left + Inches(0.3), Inches(4.5), Inches(3.2), Inches(1.2),
                    think, font_size=13, color=COLORS['text_dark'])

    # ==========================================
    # 第10页：进化视角
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, COLORS['bg_light'])

    add_rounded_rect(slide, Inches(0.5), Inches(0.4), Inches(12.3), Inches(0.8),
                     COLORS['accent'])
    add_textbox(slide, Inches(0.8), Inches(0.5), Inches(12), Inches(0.6),
                '🌱 进化视角：本阶段的"种子"', font_size=24, bold=True,
                color=COLORS['text_light'])

    seeds = [
        ('表达式可组合', '模块化组合思想', '贯穿始终', COLORS['primary']),
        ('调用printf', '接口 —— 定义契约\n隐藏实现', '第7阶段', COLORS['secondary']),
        ('标准库', '库 —— 二进制级\n复用', '第4-5阶段', COLORS['accent']),
        ('main入口', '控制反转 —— 框架\n调用你的代码', '第9阶段', COLORS['highlight']),
    ]

    for i, (seed, concept, stage, color) in enumerate(seeds):
        left = Inches(0.6 + i * 3.1)
        add_card(slide, left, Inches(1.5), Inches(2.9), Inches(2.8),
                 seed, '', title_color=color)
        add_textbox(slide, left + Inches(0.2), Inches(2.2), Inches(2.5), Inches(0.3),
                    '→ 未来概念', font_size=11, bold=True, color=COLORS['text_gray'])
        add_textbox(slide, left + Inches(0.2), Inches(2.5), Inches(2.5), Inches(1),
                    concept, font_size=13, color=COLORS['text_dark'])
        add_textbox(slide, left + Inches(0.2), Inches(3.5), Inches(2.5), Inches(0.4),
                    f'出现：{stage}', font_size=11, bold=True, color=color)

    # 痛点 & 下一阶段
    add_rounded_rect(slide, Inches(0.8), Inches(4.8), Inches(5.5), Inches(2),
                     RGBColor(0xFF, 0xEB, 0xEE), RGBColor(0xEF, 0x53, 0x50))
    add_textbox(slide, Inches(1.1), Inches(5), Inches(5), Inches(0.4),
                '😫 本阶段的痛点', font_size=16, bold=True,
                color=COLORS['danger'])
    add_textbox(slide, Inches(1.1), Inches(5.5), Inches(5), Inches(1.2),
                '所有代码都写在一个main函数里。\n'
                '代码多了之后，\n'
                'main函数会变成"屎山"。',
                font_size=13, color=COLORS['text_dark'])

    add_rounded_rect(slide, Inches(7), Inches(4.8), Inches(5.5), Inches(2),
                     RGBColor(0xE8, 0xF5, 0xE9), RGBColor(0x4C, 0xAF, 0x50))
    add_textbox(slide, Inches(7.3), Inches(5), Inches(5), Inches(0.4),
                '➡️ 下一阶段', font_size=16, bold=True,
                color=COLORS['success'])
    add_textbox(slide, Inches(7.3), Inches(5.5), Inches(5), Inches(1.2),
                '02_函数封装\n\n'
                '把代码拆分成函数，\n'
                '让main函数保持清爽。',
                font_size=13, color=COLORS['text_dark'])

    # ==========================================
    # 第11页：小结
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, COLORS['bg_light'])

    add_rounded_rect(slide, Inches(0.5), Inches(0.4), Inches(12.3), Inches(0.8),
                     COLORS['primary'])
    add_textbox(slide, Inches(0.8), Inches(0.5), Inches(12), Inches(0.6),
                '📝 本讲小结', font_size=24, bold=True,
                color=COLORS['text_light'])

    summary_data = [
        ['知识点', '要点'],
        ['C语言历史', '1972年，Dennis Ritchie，贝尔实验室，UNIX的副产品'],
        ['表达式', 'C语言的灵魂 —— 几乎一切都是表达式，都有值'],
        ['程序层次', '表达式 → 语句 → 函数 → 程序（从小到大）'],
        ['编译四步', '预处理 → 编译 → 汇编 → 链接'],
        ['接口思想', '调用者不需要知道实现，只需要知道怎么用'],
        ['进化思维', '每一个技术都是为了解决前一阶段的痛点'],
    ]
    add_table(slide, Inches(1.5), Inches(1.5), Inches(10.3), Inches(4),
              summary_data, header_color=COLORS['primary'])

    # 课后任务
    add_rounded_rect(slide, Inches(1.5), Inches(5.8), Inches(10.3), Inches(1.3),
                     RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xE0, 0xE0, 0xE0))
    add_textbox(slide, Inches(1.8), Inches(5.9), Inches(9.7), Inches(0.4),
                '📚 课后任务', font_size=16, bold=True, color=COLORS['primary'])
    add_textbox(slide, Inches(1.8), Inches(6.3), Inches(9.7), Inches(0.7),
                '① 必做：完成3个实验，记录错误和修复过程\n'
                '② 选做：用AI查"Dennis Ritchie 和 C语言的故事"，写300字感想\n'
                '③ 预习：main函数里有100行代码，会有什么问题？怎么解决？',
                font_size=12, color=COLORS['text_dark'])

    # ==========================================
    # 第12页：思考题
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, COLORS['bg_light'])

    add_rounded_rect(slide, Inches(0.5), Inches(0.4), Inches(12.3), Inches(0.8),
                     COLORS['highlight'])
    add_textbox(slide, Inches(0.8), Inches(0.5), Inches(12), Inches(0.6),
                '🤔 思辨思考题（4道）', font_size=24, bold=True,
                color=COLORS['text_light'])

    questions = [
        ('Q1', 'C语言的"命"\nC语言50年长盛不衰，\n根本原因是什么？\n会被取代吗？'),
        ('Q2', '表达式的"哲学"\n"万物皆表达式"的设计，\n优点和缺点分别是什么？'),
        ('Q3', '接口的起点\n从printf看接口思想，\n再举3个生活中的例子。'),
        ('Q4', '从表达式到框架\n大胆预测：中间需要经历\n哪些关键的"跃升"？'),
    ]

    for i, (qnum, qtext) in enumerate(questions):
        left = Inches(0.6 + i * 3.1)
        add_card(slide, left, Inches(1.5), Inches(2.9), Inches(3.5),
                 qnum, qtext, title_color=COLORS['highlight'])

    # 底部
    add_rounded_rect(slide, Inches(2), Inches(5.5), Inches(9.3), Inches(1.3),
                     COLORS['bg_dark'])
    add_textbox(slide, Inches(2.3), Inches(5.7), Inches(8.7), Inches(0.9),
                '完整题目与参考思路详见「思考题.md」\n'
                '没有标准答案，重在思考过程 💡',
                font_size=16, color=COLORS['text_light'], align=PP_ALIGN.CENTER)

    # ==========================================
    # 第13页：谢谢
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, COLORS['bg_dark'])

    # 装饰
    for i in range(6):
        box = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(1 + i * 2), Inches(6),
            Inches(0.4), Inches(0.4)
        )
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(0x2A, 0x3A, 0x50)
        box.line.color.rgb = RGBColor(0x3D, 0x5A, 0x80)

    add_textbox(slide, Inches(2), Inches(2.5), Inches(9.3), Inches(1.2),
                '谢谢！', font_size=54, bold=True,
                color=COLORS['text_light'], align=PP_ALIGN.CENTER)

    add_textbox(slide, Inches(2), Inches(3.8), Inches(9.3), Inches(0.8),
                '下一讲：函数封装 —— 告别"屎山"main',
                font_size=22, color=COLORS['secondary'], align=PP_ALIGN.CENTER)

    add_textbox(slide, Inches(2), Inches(5), Inches(9.3), Inches(0.6),
                '从程序员到软件框架 · 第01讲',
                font_size=14, color=COLORS['text_gray'], align=PP_ALIGN.CENTER)

    # 保存
    output_path = r'e:\2026-2027\2026-2027-1\AI化教学创新\AI化教学创新03\01_你好世界\docs\课件.pptx'
    prs.save(output_path)
    print(f'PPT 已生成：{output_path}')
    print(f'共 {len(prs.slides)} 页')


if __name__ == '__main__':
    create_ppt()
