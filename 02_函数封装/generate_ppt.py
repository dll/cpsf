# -*- coding: utf-8 -*-
"""
第2阶段：函数封装 - PPT生成脚本
风格：轻松愉快、卡通风、明亮配色
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# 配色方案（轻松活泼风）
COLORS = {
    'primary': RGBColor(0xFF, 0x8C, 0x42),    # 暖橙色
    'secondary': RGBColor(0x4E, 0xCD, 0xC4),     # 薄荷绿
    'accent': RGBColor(0xFF, 0x6B, 0x6B),      # 珊瑚红
    'dark': RGBColor(0x2C, 0x3E, 0x50),        # 深灰蓝
    'light': RGBColor(0xF7, 0xF9, 0xFC),       # 浅灰蓝
    'yellow': RGBColor(0xFF, 0xD9, 0x3D),      # 明黄色
    'purple': RGBColor(0xA2, 0x9B, 0xFE),      # 薰衣草紫
    'pink': RGBColor(0xFF, 0x85, 0xA6),        # 粉红色
    'green': RGBColor(0x6B, 0xCB, 0x77),       # 草绿色
    'white': RGBColor(0xFF, 0xFF, 0xFF),
    'text': RGBColor(0x2C, 0x3E, 0x50),
    'text_light': RGBColor(0x7F, 0x8C, 0x8D),
}

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

def add_bg(slide, color=COLORS['light']):
    """添加背景色"""
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg

def add_rounded_rect(slide, left, top, width, height, fill_color, border_color=None):
    """添加圆角矩形"""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(2)
    else:
        shape.line.fill.background()
    return shape

def add_text(slide, left, top, width, height, text, font_size=24, color=COLORS['text'], bold=False, align=PP_ALIGN.LEFT):
    """添加文本框"""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = '微软雅黑'
    return txBox

def add_title_bar(slide, title, subtitle=None):
    """添加标题栏"""
    # 顶部装饰条
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.15))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLORS['primary']
    bar.line.fill.background()
    
    # 标题
    add_text(slide, Inches(0.8), Inches(0.5), Inches(12), Inches(0.8), 
             title, font_size=36, color=COLORS['dark'], bold=True)
    
    if subtitle:
        add_text(slide, Inches(0.8), Inches(1.2), Inches(12), Inches(0.5),
                 subtitle, font_size=18, color=COLORS['text_light'])
    
    # 底部装饰
    add_text(slide, Inches(0.8), Inches(7.0), Inches(12), Inches(0.3),
             "从程序员到架构师 · 第2讲：函数封装", font_size=12, color=COLORS['text_light'])

# ========== 第1页：封面 ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, COLORS['light'])

# 装饰圆形
for i, (x, y, size, color) in enumerate([
    (Inches(11), Inches(0.5), Inches(2), COLORS['primary']),
    (Inches(0.5), Inches(5.5), Inches(1.5), COLORS['secondary']),
    (Inches(12), Inches(5), Inches(1), COLORS['yellow']),
    (Inches(1.5), Inches(1), Inches(0.8), COLORS['pink']),
]):
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, size, size)
    circle.fill.solid()
    circle.fill.fore_color.rgb = color
    circle.fill.transparency = 0.3
    circle.line.fill.background()

# 主标题
add_text(slide, Inches(1), Inches(2.2), Inches(11), Inches(1.2),
         "第2讲：函数封装", font_size=54, color=COLORS['dark'], bold=True, align=PP_ALIGN.CENTER)

# 副标题
add_text(slide, Inches(1), Inches(3.4), Inches(11), Inches(0.8),
         "—— 给代码找个家", font_size=32, color=COLORS['primary'], bold=True, align=PP_ALIGN.CENTER)

# 分隔线
line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5), Inches(4.3), Inches(3.3), Inches(0.05))
line.fill.solid()
line.fill.fore_color.rgb = COLORS['secondary']
line.line.fill.background()

# 系列名称
add_text(slide, Inches(1), Inches(4.7), Inches(11), Inches(0.5),
         "从程序员到架构师 · C语言插件框架演进之旅", font_size=20, color=COLORS['text_light'], align=PP_ALIGN.CENTER)

# 阶段标识
add_text(slide, Inches(5), Inches(5.5), Inches(3.3), Inches(0.6),
         "02 / 12", font_size=24, color=COLORS['accent'], bold=True, align=PP_ALIGN.CENTER)

# ========== 第2页：导入 - 屎山是怎么炼成的 ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "💩 导入：屎山是怎么炼成的？", "一个main函数的黑化之路")

# 左边：代码示例
code_box = add_rounded_rect(slide, Inches(0.5), Inches(1.8), Inches(6), Inches(5), 
                            RGBColor(0x2C, 0x3E, 0x50))
add_text(slide, Inches(0.8), Inches(2.0), Inches(5.5), Inches(0.5),
         "int main() {", font_size=18, color=COLORS['green'], bold=True)

code_lines = [
    "    // 初始化系统...  50行",
    "    // 用户登录...    80行",
    "    // 显示菜单...    60行",
    "    // 存款操作...    50行",
    "    // 取款操作...    50行",
    "    // 转账操作...    70行",
    "    // 改密码...      60行",
    "    // ...还有很多...",
    "    // 总共 800+ 行 😱",
]
for i, line in enumerate(code_lines):
    add_text(slide, Inches(0.8), Inches(2.5 + i * 0.45), Inches(5.5), Inches(0.4),
             line, font_size=14, color=RGBColor(0xBD, 0xC3, 0xC7))

add_text(slide, Inches(0.8), Inches(6.3), Inches(5.5), Inches(0.4),
         "}", font_size=18, color=COLORS['green'], bold=True)

# 右边：症状列表
add_text(slide, Inches(7), Inches(2.0), Inches(5.8), Inches(0.5),
         "😵 屎山症状", font_size=24, color=COLORS['accent'], bold=True)

symptoms = [
    ("📏", "main函数越来越长", "从10行→100行→1000行"),
    ("💥", "改一处崩十处", "牵一发而动全身"),
    ("🤔", "看不懂自己写的", "三个月后：这是我写的？"),
    ("🙈", "新人不敢碰", "看了代码直摇头"),
]

for i, (emoji, title, desc) in enumerate(symptoms):
    y = Inches(2.7 + i * 1.0)
    add_rounded_rect(slide, Inches(7), y, Inches(5.8), Inches(0.85),
                     COLORS['white'], COLORS['yellow'])
    add_text(slide, Inches(7.2), y + Inches(0.1), Inches(0.6), Inches(0.6),
             emoji, font_size=28, align=PP_ALIGN.CENTER)
    add_text(slide, Inches(7.9), y + Inches(0.08), Inches(4.7), Inches(0.4),
             title, font_size=18, color=COLORS['dark'], bold=True)
    add_text(slide, Inches(7.9), y + Inches(0.45), Inches(4.7), Inches(0.35),
             desc, font_size=13, color=COLORS['text_light'])

# ========== 第3页：什么是函数 ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "📦 什么是函数？", "一段有名字的、可以重复使用的代码块")

# 函数四件套图示
add_text(slide, Inches(1), Inches(2.0), Inches(11), Inches(0.5),
         "函数的「四件套」", font_size=28, color=COLORS['dark'], bold=True, align=PP_ALIGN.CENTER)

# 函数原型大展示
func_box = add_rounded_rect(slide, Inches(2.5), Inches(2.8), Inches(8.3), Inches(1.5),
                            COLORS['white'], COLORS['primary'])

add_text(slide, Inches(2.8), Inches(3.1), Inches(2), Inches(0.5),
         "double", font_size=28, color=COLORS['accent'], bold=True)
add_text(slide, Inches(2.8), Inches(3.65), Inches(2), Inches(0.4),
         "返回值类型", font_size=14, color=COLORS['text_light'], align=PP_ALIGN.CENTER)

add_text(slide, Inches(4.6), Inches(3.1), Inches(2), Inches(0.5),
         "deposit", font_size=28, color=COLORS['primary'], bold=True)
add_text(slide, Inches(4.6), Inches(3.65), Inches(2), Inches(0.4),
         "函数名", font_size=14, color=COLORS['text_light'], align=PP_ALIGN.CENTER)

add_text(slide, Inches(6.5), Inches(3.1), Inches(4), Inches(0.5),
         "(double balance, double amount)", font_size=22, color=COLORS['secondary'], bold=True)
add_text(slide, Inches(6.5), Inches(3.65), Inches(4), Inches(0.4),
         "参数列表", font_size=14, color=COLORS['text_light'], align=PP_ALIGN.CENTER)

# 函数体
body_box = add_rounded_rect(slide, Inches(2.5), Inches(4.5), Inches(8.3), Inches(1.8),
                            COLORS['white'], COLORS['purple'])
add_text(slide, Inches(2.8), Inches(4.7), Inches(7.7), Inches(0.4),
         "{", font_size=24, color=COLORS['purple'], bold=True)
add_text(slide, Inches(3.2), Inches(5.1), Inches(7.3), Inches(0.4),
         "    // 函数体：具体做什么", font_size=16, color=COLORS['text'])
add_text(slide, Inches(3.2), Inches(5.5), Inches(7.3), Inches(0.4),
         "    return balance + amount;   ← 返回值", font_size=16, color=COLORS['accent'], bold=True)
add_text(slide, Inches(2.8), Inches(5.9), Inches(7.7), Inches(0.4),
         "}", font_size=24, color=COLORS['purple'], bold=True)

# ========== 第4页：封装的思想 ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "🎁 封装的核心思想", "黑盒子理论：知道怎么用，不用知道怎么做")

# 左边：黑盒子示意图
add_text(slide, Inches(0.8), Inches(2.0), Inches(6), Inches(0.5),
         "黑盒子理论", font_size=28, color=COLORS['dark'], bold=True, align=PP_ALIGN.CENTER)

# 黑盒子
box = add_rounded_rect(slide, Inches(2), Inches(3.0), Inches(3.5), Inches(2.5),
                       COLORS['dark'])
add_text(slide, Inches(2), Inches(3.3), Inches(3.5), Inches(0.5),
         "⚙️", font_size=48, align=PP_ALIGN.CENTER)
add_text(slide, Inches(2), Inches(4.0), Inches(3.5), Inches(0.5),
         "deposit 函数", font_size=22, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)
add_text(slide, Inches(2), Inches(4.6), Inches(3.5), Inches(0.6),
         "（存款机）", font_size=14, color=COLORS['secondary'], align=PP_ALIGN.CENTER)

# 输入箭头
add_text(slide, Inches(0.5), Inches(3.5), Inches(1.5), Inches(0.8),
         "💰 钱\n(参数)", font_size=16, color=COLORS['primary'], bold=True, align=PP_ALIGN.CENTER)
arrow1 = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(1.8), Inches(3.8), Inches(0.5), Inches(0.3))
arrow1.fill.solid()
arrow1.fill.fore_color.rgb = COLORS['primary']
arrow1.line.fill.background()

# 输出箭头
arrow2 = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(5.3), Inches(3.8), Inches(0.5), Inches(0.3))
arrow2.fill.solid()
arrow2.fill.fore_color.rgb = COLORS['green']
arrow2.line.fill.background()
add_text(slide, Inches(5.8), Inches(3.5), Inches(1.5), Inches(0.8),
         "💵 新余额\n(返回值)", font_size=16, color=COLORS['green'], bold=True, align=PP_ALIGN.CENTER)

# 右边：封装的好处
add_text(slide, Inches(7.5), Inches(2.0), Inches(5), Inches(0.5),
         "封装的好处", font_size=28, color=COLORS['dark'], bold=True)

benefits = [
    ("🧹", "隐藏复杂度", "不用管里面怎么实现"),
    ("🔌", "标准化接口", "知道输入输出就能用"),
    ("🔧", "便于修改", "内部改了，接口不变就行"),
    ("♻️", "代码复用", "一个函数，到处调用"),
]

for i, (emoji, title, desc) in enumerate(benefits):
    y = Inches(2.7 + i * 1.0)
    add_rounded_rect(slide, Inches(7.5), y, Inches(5.3), Inches(0.85),
                     COLORS['white'], COLORS['secondary'])
    add_text(slide, Inches(7.7), y + Inches(0.1), Inches(0.8), Inches(0.6),
             emoji, font_size=28, align=PP_ALIGN.CENTER)
    add_text(slide, Inches(8.6), y + Inches(0.08), Inches(4), Inches(0.4),
             title, font_size=18, color=COLORS['dark'], bold=True)
    add_text(slide, Inches(8.6), y + Inches(0.45), Inches(4), Inches(0.35),
             desc, font_size=13, color=COLORS['text_light'])

# ========== 第5页：声明 vs 定义 ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "📜 函数声明 vs 函数定义", "接口和实现的分离——封装的雏形")

# 左边：声明
left_box = add_rounded_rect(slide, Inches(0.5), Inches(1.8), Inches(6), Inches(5.2),
                            COLORS['white'], COLORS['primary'])
add_text(slide, Inches(0.5), Inches(2.0), Inches(6), Inches(0.6),
         "📋 函数声明（接口）", font_size=24, color=COLORS['primary'], bold=True, align=PP_ALIGN.CENTER)

add_text(slide, Inches(1), Inches(2.9), Inches(5), Inches(0.5),
         "double deposit(double, double);", font_size=22, color=COLORS['dark'], bold=True, align=PP_ALIGN.CENTER)

decl_points = [
    "• 告诉编译器「有这个函数」",
    "• 只说「长什么样」，不说「怎么做」",
    "• 就像菜单：告诉你菜名，没告诉你怎么做",
    "• 放在文件开头或头文件(.h)里",
]
for i, point in enumerate(decl_points):
    add_text(slide, Inches(1), Inches(3.7 + i * 0.6), Inches(5), Inches(0.5),
             point, font_size=16, color=COLORS['text'])

# 右边：定义
right_box = add_rounded_rect(slide, Inches(6.8), Inches(1.8), Inches(6), Inches(5.2),
                             COLORS['white'], COLORS['secondary'])
add_text(slide, Inches(6.8), Inches(2.0), Inches(6), Inches(0.6),
         "🔧 函数定义（实现）", font_size=24, color=COLORS['secondary'], bold=True, align=PP_ALIGN.CENTER)

def_code = [
    "double deposit(double bal, double amt) {",
    "    if (amt > 0) {",
    '        printf("存入%.2f元\\n", amt);',
    "        return bal + amt;",
    "    }",
    "    return bal;",
    "}",
]
for i, line in enumerate(def_code):
    add_text(slide, Inches(7.2), Inches(2.9 + i * 0.45), Inches(5.3), Inches(0.4),
             line, font_size=15, color=COLORS['dark'])

def_points = [
    "• 函数的具体实现代码",
    "• 告诉电脑「具体怎么做」",
    "• 就像菜谱：详细的做法步骤",
    "• 放在.c文件里",
]
for i, point in enumerate(def_points):
    add_text(slide, Inches(7.2), Inches(6.2 + i * 0.25), Inches(5.3), Inches(0.3),
             point, font_size=13, color=COLORS['text_light'])

# ========== 第6页：三个三分之一 - AI辅助 ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "🤖 AI辅助学习", "让AI当你的编程助教")

ai_uses = [
    ("💡", "理解概念", "让AI用通俗的话解释函数封装", 
     "函数封装就像洗衣机：放脏衣服进去，按按钮，\n出来干净衣服。不用管里面怎么转的。"),
    ("✍️", "生成函数", "描述功能，让AI帮你写函数骨架", 
     "写一个判断质数的函数，参数是int，返回int。"),
    ("🔄", "重构代码", "把屎山代码丢给AI，让它拆分成函数", 
     "帮我把这段main函数拆分成多个函数。"),
    ("🐛", "调试帮忙", "函数有bug？让AI帮你分析", 
     "这个函数为什么返回值不对？帮我看看。"),
]

for i, (emoji, title, desc, example) in enumerate(ai_uses):
    col = i % 2
    row = i // 2
    x = Inches(0.5 + col * 6.3)
    y = Inches(1.8 + row * 2.7)
    
    add_rounded_rect(slide, x, y, Inches(6), Inches(2.4),
                     COLORS['white'], COLORS['purple'])
    add_text(slide, x + Inches(0.3), y + Inches(0.2), Inches(0.8), Inches(0.8),
             emoji, font_size=36, align=PP_ALIGN.CENTER)
    add_text(slide, x + Inches(1.2), y + Inches(0.25), Inches(4.5), Inches(0.5),
             title, font_size=20, color=COLORS['dark'], bold=True)
    add_text(slide, x + Inches(1.2), y + Inches(0.75), Inches(4.5), Inches(0.4),
             desc, font_size=13, color=COLORS['text_light'])
    
    # 示例框
    example_box = add_rounded_rect(slide, x + Inches(0.3), y + Inches(1.3), Inches(5.4), Inches(0.9),
                                   COLORS['light'], COLORS['purple'])
    add_text(slide, x + Inches(0.5), y + Inches(1.4), Inches(5), Inches(0.7),
             example, font_size=12, color=COLORS['text'])

# ========== 第7页：进化视角 ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "🌱 进化视角", "从函数到模块，封装的层层升级")

# 演进路线图
stages = [
    ("01", "表达式", COLORS['yellow']),
    ("02", "函数封装", COLORS['primary']),
    ("03", "模块化", COLORS['secondary']),
    ("04", "静态库", COLORS['purple']),
    ("05", "动态库", COLORS['pink']),
    ("...", "...", COLORS['text_light']),
    ("12", "插件框架", COLORS['accent']),
]

# 阶段圆圈
for i, (num, name, color) in enumerate(stages):
    x = Inches(0.5 + i * 1.75)
    y = Inches(3.0)
    
    # 连接线
    if i < len(stages) - 1:
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 
                                       x + Inches(1.1), y + Inches(0.45), 
                                       Inches(0.8), Inches(0.1))
        line.fill.solid()
        line.fill.fore_color.rgb = COLORS['text_light']
        line.line.fill.background()
    
    # 圆圈
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, Inches(1), Inches(1))
    circle.fill.solid()
    circle.fill.fore_color.rgb = color
    circle.line.fill.background()
    
    # 阶段号
    add_text(slide, x, y + Inches(0.2), Inches(1), Inches(0.4),
             num, font_size=20, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)
    
    # 阶段名
    add_text(slide, x - Inches(0.2), y + Inches(1.1), Inches(1.4), Inches(0.4),
             name, font_size=14, color=COLORS['dark'], bold=True, align=PP_ALIGN.CENTER)

# 当前位置标记
add_text(slide, Inches(2.0), Inches(2.3), Inches(1), Inches(0.5),
         "👈 当前", font_size=16, color=COLORS['primary'], bold=True, align=PP_ALIGN.CENTER)

# 下面的说明
add_text(slide, Inches(1), Inches(4.8), Inches(11), Inches(0.5),
         "函数解决了什么问题？", font_size=22, color=COLORS['dark'], bold=True, align=PP_ALIGN.CENTER)

pro_cons = [
    ("✅ 代码复用", "✅ 逻辑清晰", "✅ 便于维护", "✅ 单一职责"),
]
for i, item in enumerate(pro_cons[0]):
    x = Inches(1 + i * 2.8)
    add_rounded_rect(slide, x, Inches(5.4), Inches(2.5), Inches(0.7),
                     COLORS['green'], COLORS['green'])
    add_text(slide, x, Inches(5.5), Inches(2.5), Inches(0.5),
             item, font_size=16, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)

add_text(slide, Inches(1), Inches(6.3), Inches(11), Inches(0.5),
         "函数带来了什么新问题？→ 所有函数还在一个文件里！", font_size=20, color=COLORS['accent'], bold=True, align=PP_ALIGN.CENTER)

# ========== 第8页：小结 ==========
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "📝 小结", "函数封装，就是给代码找个家")

# 核心知识点
add_text(slide, Inches(1), Inches(2.0), Inches(11), Inches(0.6),
         "🎯 一句话总结", font_size=32, color=COLORS['dark'], bold=True, align=PP_ALIGN.CENTER)

summary_box = add_rounded_rect(slide, Inches(2), Inches(2.8), Inches(9.3), Inches(1.2),
                               COLORS['primary'], COLORS['primary'])
add_text(slide, Inches(2), Inches(3.1), Inches(9.3), Inches(0.6),
         "函数封装，就是给代码找个家——", font_size=24, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)
add_text(slide, Inches(2), Inches(3.6), Inches(9.3), Inches(0.5),
         "每个功能住一个房间，整整齐齐，互不打扰。", font_size=20, color=COLORS['white'], align=PP_ALIGN.CENTER)

# 五个核心点
key_points = [
    ("1️⃣", "函数是什么", "有名字的可复用代码块"),
    ("2️⃣", "函数四件套", "返回值、函数名、参数、函数体"),
    ("3️⃣", "封装思想", "隐藏实现，只露接口"),
    ("4️⃣", "声明vs定义", "接口（.h）和实现（.c）"),
    ("5️⃣", "函数的好处", "复用、清晰、易调试、易协作"),
]

for i, (num, title, desc) in enumerate(key_points):
    x = Inches(0.5 + i * 2.55)
    y = Inches(4.4)
    
    add_rounded_rect(slide, x, y, Inches(2.35), Inches(2.2),
                     COLORS['white'], COLORS['secondary'])
    add_text(slide, x, y + Inches(0.15), Inches(2.35), Inches(0.5),
             num, font_size=28, align=PP_ALIGN.CENTER)
    add_text(slide, x, y + Inches(0.7), Inches(2.35), Inches(0.4),
             title, font_size=16, color=COLORS['dark'], bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, x + Inches(0.15), y + Inches(1.2), Inches(2.05), Inches(0.8),
             desc, font_size=13, color=COLORS['text_light'], align=PP_ALIGN.CENTER)

# 保存
output_path = r"e:\2026-2027\2026-2027-1\AI化教学创新\AI化教学创新03\02_函数封装\docs\课件.pptx"
prs.save(output_path)
print(f"PPT生成完成：{output_path}")
