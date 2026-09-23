# -*- coding: utf-8 -*-
"""
第3阶段：函数封装 —— PPT生成脚本
风格：轻松愉快、卡通风、明亮配色
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path

COLORS = {
    'primary': RGBColor(0xFF, 0x8C, 0x42),
    'secondary': RGBColor(0x4E, 0xCD, 0xC4),
    'accent': RGBColor(0xFF, 0x6B, 0x6B),
    'dark': RGBColor(0x2C, 0x3E, 0x50),
    'light': RGBColor(0xF7, 0xF9, 0xFC),
    'yellow': RGBColor(0xFF, 0xD9, 0x3D),
    'purple': RGBColor(0xA2, 0x9B, 0xFE),
    'pink': RGBColor(0xFF, 0x85, 0xA6),
    'green': RGBColor(0x6B, 0xCB, 0x77),
    'blue': RGBColor(0x74, 0xB9, 0xFF),
    'white': RGBColor(0xFF, 0xFF, 0xFF),
    'text': RGBColor(0x2C, 0x3E, 0x50),
    'text_light': RGBColor(0x7F, 0x8C, 0x8D),
}

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

def add_bg(slide, color=COLORS['light']):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg

def add_rounded_rect(slide, left, top, width, height, fill_color, border_color=None, radius=0.08):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(2)
    else:
        shape.line.fill.background()
    shape.adjustments[0] = radius
    return shape

def add_circle(slide, left, top, size, fill_color):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    return shape

def add_text(slide, left, top, width, height, text, font_size=24, color=COLORS['text'],
             bold=False, align=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = '微软雅黑'
    return txBox

def add_title_bar(slide, title, subtitle=None):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLORS['primary']
    bar.line.fill.background()
    add_circle(slide, Inches(0.3), Inches(0.5), Inches(0.35), COLORS['primary'])
    add_circle(slide, Inches(0.55), Inches(0.55), Inches(0.25), COLORS['secondary'])
    add_text(slide, Inches(1.0), Inches(0.4), Inches(11), Inches(0.7),
             title, font_size=32, color=COLORS['dark'], bold=True)
    if subtitle:
        add_text(slide, Inches(1.0), Inches(1.05), Inches(11), Inches(0.45),
                 subtitle, font_size=15, color=COLORS['text_light'])
    add_text(slide, Inches(0.8), Inches(7.05), Inches(12), Inches(0.3),
             "第3讲 函数封装", font_size=11, color=COLORS['text_light'])

def add_shape(slide, shape_type, left, top, width, height, fill_color=None, line_color=None):
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    if fill_color is not None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if line_color is not None:
        shape.line.color.rgb = line_color
    else:
        shape.line.fill.background()
    return shape

# ============================================================
# 第1页：封面
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, COLORS['light'])
add_circle(slide, Inches(9.5), Inches(-1.5), Inches(5), COLORS['secondary'])
add_circle(slide, Inches(-2), Inches(4.5), Inches(4.5), COLORS['primary'])
add_circle(slide, Inches(11), Inches(5), Inches(2.5), COLORS['yellow'])

card = add_rounded_rect(slide, Inches(1.5), Inches(1.5), Inches(10.3), Inches(4.5),
                        COLORS['white'], COLORS['white'], 0.05)

# 左侧 - 工具箱图标
box_x = Inches(2.5)
box_y = Inches(2.3)
add_rounded_rect(slide, box_x, box_y, Inches(2.2), Inches(2), COLORS['accent'], COLORS['accent'], 0.1)
add_rounded_rect(slide, box_x + Inches(0.3), box_y - Inches(0.3), Inches(1.6), Inches(0.5), COLORS['accent'], COLORS['accent'], 0.2)
add_text(slide, box_x, box_y + Inches(0.4), Inches(2.2), Inches(1.2), "🔧", font_size=56, align=PP_ALIGN.CENTER)

# 右侧文字
add_text(slide, Inches(5.2), Inches(2.3), Inches(6), Inches(1),
         "第3讲：函数封装", font_size=52, color=COLORS['dark'], bold=True)
add_text(slide, Inches(5.2), Inches(3.3), Inches(6), Inches(0.6),
         "给代码找个家", font_size=26, color=COLORS['primary'], bold=True)
divider = add_shape(slide, MSO_SHAPE.RECTANGLE, Inches(5.2), Inches(4.1), Inches(1.5), Inches(0.06), COLORS['secondary'])
add_text(slide, Inches(5.2), Inches(4.3), Inches(6), Inches(0.5),
         "《从程序员到架构师》· C语言插件框架演进之旅", font_size=16, color=COLORS['text_light'])
add_text(slide, Inches(5.2), Inches(4.8), Inches(6), Inches(0.5),
         "封装 · 复用 · 单一职责", font_size=18, color=COLORS['accent'], bold=True)
add_text(slide, Inches(5.5), Inches(5.4), Inches(2.3), Inches(0.5),
         "03 / 14", font_size=22, color=COLORS['accent'], bold=True, align=PP_ALIGN.CENTER)

# ============================================================
# 第2页：知识图谱
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "🗺️ 知识图谱 · 第3讲", "本讲在整个知识体系中的位置")

center_x = Inches(5.8)
center_y = Inches(3.7)
add_circle(slide, center_x - Inches(1.5), center_y - Inches(1), Inches(3), COLORS['secondary'])
add_circle(slide, center_x - Inches(1.1), center_y - Inches(0.6), Inches(2.2), COLORS['secondary'])

center_card = add_rounded_rect(slide, center_x - Inches(1.3), center_y - Inches(0.55), Inches(2.6), Inches(1.1),
                               COLORS['secondary'], COLORS['secondary'], 0.3)
add_text(slide, center_x - Inches(1.3), center_y - Inches(0.45), Inches(2.6), Inches(0.45),
         "🔧 函数封装", font_size=24, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)
add_text(slide, center_x - Inches(1.3), center_y + Inches(0.1), Inches(2.6), Inches(0.35),
         "（第3讲 核心）", font_size=13, color=COLORS['white'], align=PP_ALIGN.CENTER)

nodes = [
    ("控制结构", COLORS['primary'], Inches(1.8), Inches(1.3), "第2讲 · 前置", "🔀"),
    ("表达式", COLORS['yellow'], Inches(0.5), Inches(3.8), "第1讲 · 基础", "📌"),
    ("多文件编程", COLORS['purple'], Inches(9.8), Inches(1.3), "→ 第4讲", "📂"),
    ("指针", COLORS['blue'], Inches(9.8), Inches(3.8), "→ 第5讲", "🎯"),
    ("ATM案例", COLORS['green'], Inches(5.3), Inches(6.0), "贯穿全系列", "🏧"),
]

for name, color, x, y, desc, icon in nodes:
    add_rounded_rect(slide, x, y, Inches(2.4), Inches(1.0), COLORS['white'], color, 0.15)
    add_text(slide, x + Inches(0.1), y + Inches(0.15), Inches(0.6), Inches(0.7),
             icon, font_size=28, align=PP_ALIGN.CENTER)
    add_text(slide, x + Inches(0.7), y + Inches(0.1), Inches(1.6), Inches(0.4),
             name, font_size=15, color=COLORS['dark'], bold=True)
    add_text(slide, x + Inches(0.7), y + Inches(0.5), Inches(1.6), Inches(0.4),
             desc, font_size=10, color=COLORS['text_light'])

connections = [
    (center_x - Inches(0.5), center_y, Inches(3.0), Inches(1.8)),
    (center_x - Inches(0.5), center_y + Inches(0.2), Inches(2.0), Inches(4.3)),
    (center_x + Inches(0.5), center_y, Inches(9.8), Inches(1.8)),
    (center_x + Inches(0.5), center_y + Inches(0.2), Inches(9.8), Inches(4.3)),
    (center_x, center_y + Inches(0.5), Inches(6.5), Inches(6.5)),
]
for sx, sy, ex, ey in connections:
    line = slide.shapes.add_connector(1, sx, sy, ex, ey)
    line.line.color.rgb = COLORS['text_light']
    line.line.width = Pt(1.5)
    line.line.dash_style = 1

tip = add_rounded_rect(slide, Inches(1.5), Inches(6.5), Inches(10.3), Inches(0.6),
                      COLORS['yellow'], COLORS['yellow'], 0.3)
add_text(slide, Inches(1.5), Inches(6.58), Inches(10.3), Inches(0.45),
         "💡 函数是代码的家——把相关代码打包、命名、隐藏细节，让代码更清晰、更好复用",
         font_size=14, color=COLORS['dark'], bold=True, align=PP_ALIGN.CENTER)

# ============================================================
# 第3页：第2讲的痛——代码太乱
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "🤔 第2讲的痛", "所有代码堆在一起——又长又乱又重复")

problems = [
    ("📏", "main太长", "几十行代码堆在一起\n加功能=越来越长", COLORS['accent']),
    ("📋", "代码重复", "存款取款转账的逻辑\n几乎一样，复制粘贴", COLORS['yellow']),
    ("🔍", "阅读困难", "想改取款？\n在一大堆代码里找半天", COLORS['purple']),
    ("📦", "没法复用", "另一个程序要用存款？\n只能整段复制", COLORS['blue']),
]

for i, (icon, title, desc, color) in enumerate(problems):
    x = Inches(0.6 + i * 3.15)
    y = Inches(2.0)
    add_rounded_rect(slide, x, y, Inches(2.9), Inches(3.5), COLORS['white'], color, 0.12)
    top_bar = add_shape(slide, MSO_SHAPE.RECTANGLE, x, y, Inches(2.9), Inches(0.08), color)
    add_text(slide, x, y + Inches(0.3), Inches(2.9), Inches(1.0), icon, font_size=48, align=PP_ALIGN.CENTER)
    add_text(slide, x, y + Inches(1.4), Inches(2.9), Inches(0.5),
             title, font_size=18, color=COLORS['dark'], bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, x + Inches(0.2), y + Inches(2.0), Inches(2.5), Inches(1.2),
             desc, font_size=13, color=COLORS['text_light'], align=PP_ALIGN.CENTER)

tip = add_rounded_rect(slide, Inches(2), Inches(5.9), Inches(9.3), Inches(1.2),
                      COLORS['primary'], COLORS['primary'], 0.3)
add_text(slide, Inches(2), Inches(6.0), Inches(9.3), Inches(0.5),
         "💡 怎么办？给代码找个家——函数封装！",
         font_size=20, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)
add_text(slide, Inches(2), Inches(6.5), Inches(9.3), Inches(0.5),
         "把相关代码打包在一起，起个名字，各归各位",
         font_size=16, color=COLORS['yellow'], bold=True, align=PP_ALIGN.CENTER)

# ============================================================
# 第4页：函数就是抽屉
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "🗄️ 函数就是抽屉", "把零散的东西分类放进抽屉——整洁又好找")

# 左侧：没有抽屉的桌面（乱）
left_x = Inches(0.5)
add_rounded_rect(slide, left_x, Inches(2.0), Inches(5.5), Inches(4.5), COLORS['white'], COLORS['accent'], 0.05)
add_text(slide, left_x, Inches(2.1), Inches(5.5), Inches(0.5),
         "❌ 没有函数（乱的）", font_size=18, color=COLORS['accent'], bold=True, align=PP_ALIGN.CENTER)
# 模拟乱代码
messy = ["printf(...)", "scanf(...)", "if (amount>0)", "balance+=amount", "printf(...)", "scanf(...)", "if (amount>bal)", "balance-=amount"]
for i, code in enumerate(messy):
    add_text(slide, left_x + Inches(0.3 + (i%2)*2.5), Inches(2.8 + (i//2)*0.55), Inches(2.3), Inches(0.4),
             code, font_size=12, color=COLORS['text_light'])

# 右侧：有抽屉的桌面（整齐）
right_x = Inches(6.5)
add_rounded_rect(slide, right_x, Inches(2.0), Inches(6.3), Inches(4.5), COLORS['white'], COLORS['green'], 0.05)
add_text(slide, right_x, Inches(2.1), Inches(6.3), Inches(0.5),
         "✅ 有了函数（整齐）", font_size=18, color=COLORS['green'], bold=True, align=PP_ALIGN.CENTER)

drawers = [
    ("deposit()", "存款逻辑", COLORS['blue']),
    ("withdraw()", "取款逻辑", COLORS['purple']),
    ("transfer()", "转账逻辑", COLORS['yellow']),
    ("print_success()", "成功提示", COLORS['green']),
]
for i, (name, desc, color) in enumerate(drawers):
    dx = right_x + Inches(0.3 + (i%2)*3.0)
    dy = Inches(2.8 + (i//2)*1.5)
    add_rounded_rect(slide, dx, dy, Inches(2.7), Inches(1.2), color, color, 0.15)
    add_text(slide, dx, dy + Inches(0.1), Inches(2.7), Inches(0.5),
             name, font_size=14, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, dx, dy + Inches(0.6), Inches(2.7), Inches(0.4),
             desc, font_size=11, color=COLORS['white'], align=PP_ALIGN.CENTER)

# 中间箭头
arrow = add_shape(slide, MSO_SHAPE.RIGHT_ARROW, Inches(5.8), Inches(4.0), Inches(0.8), Inches(0.5), COLORS['primary'])

# ============================================================
# 第5页：函数语法三阶段
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "📝 函数语法：三阶段", "声明 → 定义 → 调用（菜单 → 厨房 → 点菜）")

stages = [
    ("📋 声明", "函数原型", "void deposit(void);", "告诉编译器\n有这个函数", COLORS['blue'], "就像菜单：\n告诉你有什么菜"),
    ("🍳 定义", "函数实现", "void deposit() {\n  // 具体代码\n}", "函数的具体实现\n怎么做这件事", COLORS['green'], "就像厨房：\n真正做菜的地方"),
    ("🔔 调用", "使用函数", "deposit();", "让函数执行一次", COLORS['accent'], "就像点菜：\n告诉厨房我要这个"),
]

for i, (icon, name, code, desc, color, analogy) in enumerate(stages):
    x = Inches(0.5 + i * 4.3)
    y = Inches(2.0)
    add_rounded_rect(slide, x, y, Inches(4.0), Inches(4.5), COLORS['white'], color, 0.12)
    top_bar = add_shape(slide, MSO_SHAPE.RECTANGLE, x, y, Inches(4.0), Inches(0.7), color)
    add_text(slide, x, y + Inches(0.1), Inches(4.0), Inches(0.5),
             icon + " " + name, font_size=18, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)
    
    code_card = add_rounded_rect(slide, x + Inches(0.3), y + Inches(0.9), Inches(3.4), Inches(1.2),
                                 RGBColor(0x1E, 0x29, 0x3B), RGBColor(0x1E, 0x29, 0x3B), 0.08)
    add_text(slide, x + Inches(0.5), y + Inches(1.0), Inches(3.0), Inches(1.0),
             code, font_size=13, color=COLORS['green'])
    
    add_text(slide, x + Inches(0.3), y + Inches(2.3), Inches(3.4), Inches(0.8),
             desc, font_size=13, color=COLORS['text'])
    
    add_rounded_rect(slide, x + Inches(0.3), y + Inches(3.3), Inches(3.4), Inches(0.9),
                     COLORS['light'], color, 0.2)
    add_text(slide, x + Inches(0.3), y + Inches(3.4), Inches(3.4), Inches(0.7),
             analogy, font_size=12, color=COLORS['text_light'], align=PP_ALIGN.CENTER)

# 箭头连接
for i in range(2):
    x = Inches(4.3 + i * 4.3)
    add_shape(slide, MSO_SHAPE.RIGHT_ARROW, x, Inches(4.0), Inches(0.4), Inches(0.35), COLORS['primary'])

# ============================================================
# 第6页：参数传递
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "📦 参数传递", "让函数可以接收外部数据——做什么菜看食材")

# 左侧 - 代码示例
left_x = Inches(0.8)
code_card = add_rounded_rect(slide, left_x, Inches(2.0), Inches(6.0), Inches(4.5),
                             RGBColor(0x1E, 0x29, 0x3B), RGBColor(0x1E, 0x29, 0x3B), 0.08)
top_c = add_shape(slide, MSO_SHAPE.RECTANGLE, left_x, Inches(2.0), Inches(6.0), Inches(0.4), RGBColor(0x34, 0x49, 0x5E))
for i, c in enumerate([COLORS['accent'], COLORS['yellow'], COLORS['green']]):
    add_circle(slide, left_x + Inches(0.2 + i * 0.35), Inches(2.08), Inches(0.2), c)
add_text(slide, left_x + Inches(2.5), Inches(2.03), Inches(1), Inches(0.35),
         "参数传递.c", font_size=11, color=COLORS['text_light'], align=PP_ALIGN.CENTER)

code_lines = [
    ('// 不带参数：只能做固定的事', COLORS['text_light']),
    ('void say_hello() {', COLORS['blue']),
    ('    printf("Hello!\\n");', COLORS['green']),
    ('}', COLORS['blue']),
    ('', COLORS['text']),
    ('// 带参数：根据传入的值做不同的事', COLORS['text_light']),
    ('void print_success(char* op, double amt) {', COLORS['purple']),
    ('    printf("✅ %s %.2f元\\n", op, amt);', COLORS['green']),
    ('}', COLORS['purple']),
    ('', COLORS['text']),
    ('// 调用时传入不同的值', COLORS['text_light']),
    ('print_success("存入", 100.0);  // 存入', COLORS['yellow']),
    ('print_success("取出", 50.0);   // 取出', COLORS['yellow']),
]

for i, (code, color) in enumerate(code_lines):
    add_text(slide, left_x + Inches(0.3), Inches(2.55 + i * 0.30), Inches(5.4), Inches(0.28),
             code, font_size=12, color=color)

# 右侧 - 比喻图示
right_x = Inches(7.5)
add_rounded_rect(slide, right_x, Inches(2.0), Inches(5.2), Inches(4.5), COLORS['white'], COLORS['secondary'], 0.08)
add_shape(slide, MSO_SHAPE.RECTANGLE, right_x, Inches(2.0), Inches(5.2), Inches(0.55), COLORS['secondary'])
add_text(slide, right_x, Inches(2.08), Inches(5.2), Inches(0.4),
         "🍳 参数 = 食材", font_size=17, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)

add_text(slide, right_x + Inches(0.3), Inches(2.8), Inches(4.6), Inches(0.5),
         "同一个厨房（函数），不同的食材（参数）→", font_size=14, color=COLORS['text'])
add_text(slide, right_x + Inches(0.3), Inches(3.3), Inches(4.6), Inches(0.5),
         "做出不同的菜（不同的结果）", font_size=14, color=COLORS['text'], bold=True)

items = [
    ('print_success("存入", 100)', '→ "✅ 存入 100.00元"', COLORS['green']),
    ('print_success("取出", 50)', '→ "✅ 取出 50.00元"', COLORS['accent']),
]
for i, (code, result, color) in enumerate(items):
    y = Inches(4.2 + i * 1.0)
    add_rounded_rect(slide, right_x + Inches(0.3), y, Inches(4.6), Inches(0.8), COLORS['light'], color, 0.15)
    add_text(slide, right_x + Inches(0.4), y + Inches(0.05), Inches(4.4), Inches(0.35),
             code, font_size=11, color=COLORS['text'])
    add_text(slide, right_x + Inches(0.4), y + Inches(0.4), Inches(4.4), Inches(0.35),
             result, font_size=12, color=color, bold=True)

# ============================================================
# 第7页：值传递
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "📋 值传递", "传的是副本，不是原件——改副本不影响原件")

# 左侧 - 代码
left_x = Inches(0.8)
code_card = add_rounded_rect(slide, left_x, Inches(2.0), Inches(6.0), Inches(4.5),
                             RGBColor(0x1E, 0x29, 0x3B), RGBColor(0x1E, 0x29, 0x3B), 0.08)
add_shape(slide, MSO_SHAPE.RECTANGLE, left_x, Inches(2.0), Inches(6.0), Inches(0.4), RGBColor(0x34, 0x49, 0x5E))
add_text(slide, left_x + Inches(2.5), Inches(2.03), Inches(1), Inches(0.35),
         "值传递.c", font_size=11, color=COLORS['text_light'], align=PP_ALIGN.CENTER)

code_lines = [
    ('void add_one(int x) {', COLORS['blue']),
    ('    x = x + 1;  // 改的是x（副本）', COLORS['green']),
    ('    printf("函数内: x=%d\\n", x);', COLORS['green']),
    ('}', COLORS['blue']),
    ('', COLORS['text']),
    ('int main() {', COLORS['blue']),
    ('    int a = 10;', COLORS['yellow']),
    ('    add_one(a);  // 传a的值(10)的副本', COLORS['purple']),
    ('    printf("函数外: a=%d\\n", a);', COLORS['yellow']),
    ('    return 0;', COLORS['blue']),
    ('}', COLORS['blue']),
    ('', COLORS['text']),
    ('// 输出：', COLORS['text_light']),
    ('// 函数内: x=11  ← 改了副本', COLORS['accent']),
    ('// 函数外: a=10  ← 原件没变！', COLORS['accent']),
]

for i, (code, color) in enumerate(code_lines):
    add_text(slide, left_x + Inches(0.3), Inches(2.55 + i * 0.28), Inches(5.4), Inches(0.26),
             code, font_size=12, color=color)

# 右侧 - 复印件比喻
right_x = Inches(7.5)
add_rounded_rect(slide, right_x, Inches(2.0), Inches(5.2), Inches(4.5), COLORS['white'], COLORS['accent'], 0.08)
add_shape(slide, MSO_SHAPE.RECTANGLE, right_x, Inches(2.0), Inches(5.2), Inches(0.55), COLORS['accent'])
add_text(slide, right_x, Inches(2.08), Inches(5.2), Inches(0.4),
         "📄 复印件比喻", font_size=17, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)

# 原件
add_rounded_rect(slide, right_x + Inches(0.5), Inches(2.9), Inches(2), Inches(1.5), COLORS['yellow'], COLORS['yellow'], 0.1)
add_text(slide, right_x + Inches(0.5), Inches(3.0), Inches(2), Inches(0.4),
         "原件 a=10", font_size=14, color=COLORS['dark'], bold=True, align=PP_ALIGN.CENTER)
add_text(slide, right_x + Inches(0.5), Inches(3.5), Inches(2), Inches(0.5),
         "（在main里）", font_size=11, color=COLORS['text_light'], align=PP_ALIGN.CENTER)

# 箭头
arrow = add_shape(slide, MSO_SHAPE.RIGHT_ARROW, right_x + Inches(2.7), Inches(3.3), Inches(0.6), Inches(0.4), COLORS['accent'])
add_text(slide, right_x + Inches(2.6), Inches(2.9), Inches(1), Inches(0.3), "复制", font_size=10, color=COLORS['accent'], align=PP_ALIGN.CENTER)

# 副本
add_rounded_rect(slide, right_x + Inches(3.5), Inches(2.9), Inches(1.2), Inches(1.5), COLORS['light'], COLORS['accent'], 0.1)
add_text(slide, right_x + Inches(3.5), Inches(3.0), Inches(1.2), Inches(0.4),
         "副本 x", font_size=14, color=COLORS['accent'], bold=True, align=PP_ALIGN.CENTER)
add_text(slide, right_x + Inches(3.5), Inches(3.4), Inches(1.2), Inches(0.4),
         "→11", font_size=14, color=COLORS['green'], bold=True, align=PP_ALIGN.CENTER)
add_text(slide, right_x + Inches(3.5), Inches(3.8), Inches(1.2), Inches(0.4),
         "（改了它）", font_size=10, color=COLORS['text_light'], align=PP_ALIGN.CENTER)

add_rounded_rect(slide, right_x + Inches(0.5), Inches(4.8), Inches(4.2), Inches(1.4), COLORS['light'], COLORS['secondary'], 0.2)
add_text(slide, right_x + Inches(0.5), Inches(4.9), Inches(4.2), Inches(0.4),
         "💡 想改原件？用指针！", font_size=16, color=COLORS['secondary'], bold=True, align=PP_ALIGN.CENTER)
add_text(slide, right_x + Inches(0.5), Inches(5.3), Inches(4.2), Inches(0.7),
         "传变量的地址，函数就知道\n原件在哪，直接改原件\n→ 第5讲：指针", font_size=12, color=COLORS['text_light'], align=PP_ALIGN.CENTER)

# ============================================================
# 第8页：返回值
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "↩️ 返回值", "函数执行完，把结果带回来——上菜")

# 左侧 - 代码
left_x = Inches(0.8)
code_card = add_rounded_rect(slide, left_x, Inches(2.0), Inches(6.0), Inches(4.5),
                             RGBColor(0x1E, 0x29, 0x3B), RGBColor(0x1E, 0x29, 0x3B), 0.08)
add_shape(slide, MSO_SHAPE.RECTANGLE, left_x, Inches(2.0), Inches(6.0), Inches(0.4), RGBColor(0x34, 0x49, 0x5E))
add_text(slide, left_x + Inches(2.5), Inches(2.03), Inches(1), Inches(0.35),
         "返回值.c", font_size=11, color=COLORS['text_light'], align=PP_ALIGN.CENTER)

code_lines = [
    ('// 计算两数之和，返回结果', COLORS['text_light']),
    ('int add(int a, int b) {', COLORS['blue']),
    ('    return a + b;  // 返回计算结果', COLORS['green']),
    ('}', COLORS['blue']),
    ('', COLORS['text']),
    ('int main() {', COLORS['blue']),
    ('    // 调用函数，接收返回值', COLORS['text_light']),
    ('    int result = add(3, 5);', COLORS['yellow']),
    ('    printf("3+5=%d\\n", result);  // 输出8', COLORS['green']),
    ('    return 0;', COLORS['blue']),
    ('}', COLORS['blue']),
    ('', COLORS['text']),
    ('// return的作用：', COLORS['text_light']),
    ('// 1. 结束函数（跳出去）', COLORS['accent']),
    ('// 2. 返回一个值给调用者', COLORS['accent']),
]

for i, (code, color) in enumerate(code_lines):
    add_text(slide, left_x + Inches(0.3), Inches(2.55 + i * 0.28), Inches(5.4), Inches(0.26),
             code, font_size=12, color=color)

# 右侧 - 上菜比喻
right_x = Inches(7.5)
add_rounded_rect(slide, right_x, Inches(2.0), Inches(5.2), Inches(4.5), COLORS['white'], COLORS['green'], 0.08)
add_shape(slide, MSO_SHAPE.RECTANGLE, right_x, Inches(2.0), Inches(5.2), Inches(0.55), COLORS['green'])
add_text(slide, right_x, Inches(2.08), Inches(5.2), Inches(0.4),
         "🍽️ 上菜比喻", font_size=17, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)

# 厨房 → 上菜 → 顾客
add_rounded_rect(slide, right_x + Inches(0.5), Inches(3.0), Inches(1.5), Inches(1.2), COLORS['yellow'], COLORS['yellow'], 0.15)
add_text(slide, right_x + Inches(0.5), Inches(3.1), Inches(1.5), Inches(0.5), "🍳", font_size=32, align=PP_ALIGN.CENTER)
add_text(slide, right_x + Inches(0.5), Inches(3.7), Inches(1.5), Inches(0.4), "厨房", font_size=12, color=COLORS['dark'], align=PP_ALIGN.CENTER)

arrow1 = add_shape(slide, MSO_SHAPE.RIGHT_ARROW, right_x + Inches(2.1), Inches(3.2), Inches(1.0), Inches(0.5), COLORS['green'])
add_text(slide, right_x + Inches(2.1), Inches(3.7), Inches(1.0), Inches(0.3), "return", font_size=11, color=COLORS['green'], bold=True, align=PP_ALIGN.CENTER)

add_rounded_rect(slide, right_x + Inches(3.3), Inches(3.0), Inches(1.5), Inches(1.2), COLORS['secondary'], COLORS['secondary'], 0.15)
add_text(slide, right_x + Inches(3.3), Inches(3.1), Inches(1.5), Inches(0.5), "🍽️", font_size=32, align=PP_ALIGN.CENTER)
add_text(slide, right_x + Inches(3.3), Inches(3.7), Inches(1.5), Inches(0.4), "上菜", font_size=12, color=COLORS['white'], align=PP_ALIGN.CENTER)

add_rounded_rect(slide, right_x + Inches(0.5), Inches(4.6), Inches(4.2), Inches(1.6), COLORS['light'], COLORS['blue'], 0.15)
add_text(slide, right_x + Inches(0.5), Inches(4.7), Inches(4.2), Inches(0.4),
         "return 两个作用：", font_size=14, color=COLORS['blue'], bold=True, align=PP_ALIGN.CENTER)
add_text(slide, right_x + Inches(0.5), Inches(5.1), Inches(4.2), Inches(0.4),
         "① 结束函数——立刻跳出去", font_size=13, color=COLORS['text'])
add_text(slide, right_x + Inches(0.5), Inches(5.5), Inches(4.2), Inches(0.4),
         "② 返回值——把结果带给调用者", font_size=13, color=COLORS['text'])

# ============================================================
# 第9页：对比第2讲 vs 第3讲
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "📊 第2讲 vs 第3讲", '从"大杂烩"到"各归各位"——变化一目了然')

comparisons = [
    ("main长度", "几十行堆在一起", "十几行只调度", COLORS['accent'], COLORS['green']),
    ("代码重复", "到处复制粘贴", "抽成公共函数", COLORS['accent'], COLORS['green']),
    ("找代码", "在main里慢慢找", "直接找对应函数", COLORS['accent'], COLORS['green']),
    ("代码复用", "只能复制粘贴", "直接调用函数", COLORS['accent'], COLORS['green']),
    ("可读性", "从头读到尾", "看函数名就知道", COLORS['accent'], COLORS['green']),
    ("可维护", "改一处查所有副本", "改一个函数就行", COLORS['accent'], COLORS['green']),
]

for i, (item, old, new, old_color, new_color) in enumerate(comparisons):
    y = Inches(2.0 + i * 0.75)
    
    # 项目名
    add_rounded_rect(slide, Inches(0.8), y, Inches(2.5), Inches(0.6), COLORS['white'], COLORS['blue'], 0.15)
    add_text(slide, Inches(0.8), y + Inches(0.1), Inches(2.5), Inches(0.4),
             item, font_size=14, color=COLORS['dark'], bold=True, align=PP_ALIGN.CENTER)
    
    # 第2讲
    add_rounded_rect(slide, Inches(3.6), y, Inches(4.0), Inches(0.6), COLORS['light'], old_color, 0.15)
    add_text(slide, Inches(3.6), y + Inches(0.1), Inches(4.0), Inches(0.4),
             "❌ " + old, font_size=13, color=old_color, align=PP_ALIGN.CENTER)
    
    # 第3讲
    add_rounded_rect(slide, Inches(7.9), y, Inches(4.5), Inches(0.6), COLORS['light'], new_color, 0.15)
    add_text(slide, Inches(7.9), y + Inches(0.1), Inches(4.5), Inches(0.4),
             "✅ " + new, font_size=13, color=new_color, align=PP_ALIGN.CENTER)

# 标题行
add_text(slide, Inches(3.6), Inches(1.5), Inches(4.0), Inches(0.4),
         "第2讲（无函数）", font_size=15, color=COLORS['accent'], bold=True, align=PP_ALIGN.CENTER)
add_text(slide, Inches(7.9), Inches(1.5), Inches(4.5), Inches(0.4),
         "第3讲（函数封装）", font_size=15, color=COLORS['green'], bold=True, align=PP_ALIGN.CENTER)

gold = add_rounded_rect(slide, Inches(1.5), Inches(6.5), Inches(10.3), Inches(0.6),
                        COLORS['primary'], COLORS['primary'], 0.3)
add_text(slide, Inches(1.5), Inches(6.58), Inches(10.3), Inches(0.45),
         '🏆 封装的力量——让代码从"能用"变成"好用、好改、好理解"',
         font_size=14, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)

# ============================================================
# 第10页：设计原则
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "🎯 三个设计原则", "好代码的共同特征——从函数开始养成习惯")

principles = [
    ("1️⃣", "单一职责", "一个函数只做一件事\n并且把它做好", "✅ deposit() 只管存款\n❌ deposit_and_withdraw() 混在一起", COLORS['blue']),
    ("2️⃣", "代码复用", "同样的代码不要写两遍\n抽成函数，到处调用", "存款取款转账都有成功提示\n→ 抽成 print_success() 函数", COLORS['green']),
    ("3️⃣", "信息隐藏", "调用者不需要知道\n内部怎么实现的", "你用printf不需要知道它\n内部怎么把字显示到屏幕上", COLORS['purple']),
]

for i, (num, name, desc, example, color) in enumerate(principles):
    x = Inches(0.5 + i * 4.3)
    y = Inches(2.0)
    add_rounded_rect(slide, x, y, Inches(4.0), Inches(4.5), COLORS['white'], color, 0.12)
    top_bar = add_shape(slide, MSO_SHAPE.RECTANGLE, x, y, Inches(4.0), Inches(0.08), color)
    
    add_text(slide, x, y + Inches(0.2), Inches(4.0), Inches(0.6), num, font_size=32, align=PP_ALIGN.CENTER)
    add_text(slide, x, y + Inches(0.9), Inches(4.0), Inches(0.5),
             name, font_size=20, color=COLORS['dark'], bold=True, align=PP_ALIGN.CENTER)
    
    divider = add_shape(slide, MSO_SHAPE.RECTANGLE, x + Inches(1.3), y + Inches(1.5), Inches(1.4), Inches(0.04), color)
    
    add_text(slide, x + Inches(0.3), y + Inches(1.7), Inches(3.4), Inches(0.8),
             desc, font_size=13, color=COLORS['text_light'], align=PP_ALIGN.CENTER)
    
    add_rounded_rect(slide, x + Inches(0.3), y + Inches(2.7), Inches(3.4), Inches(1.5),
                     COLORS['light'], color, 0.15)
    add_text(slide, x + Inches(0.4), y + Inches(2.8), Inches(3.2), Inches(1.3),
             example, font_size=12, color=COLORS['text'])

# 底部 - 封装演进
gold = add_rounded_rect(slide, Inches(1.5), Inches(6.6), Inches(10.3), Inches(0.6),
                        COLORS['yellow'], COLORS['yellow'], 0.3)
add_text(slide, Inches(1.5), Inches(6.68), Inches(10.3), Inches(0.45),
         "💡 这三个原则贯穿整个系列：函数→模块→库→框架，都是在更高层次做同样的事",
         font_size=13, color=COLORS['dark'], bold=True, align=PP_ALIGN.CENTER)

# ============================================================
# 第11页：思考题
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "🤔 思考题 · 往深了想", "没有标准答案，重要的是思考过程")

questions = [
    ("函数=封装？", "为什么说函数是C语言中\n最基础的封装手段？\n封装到底是什么意思？", COLORS['primary'], Inches(0.8), "💡"),
    ("值传递的得与失", "值传递有什么好处？\n有什么坏处？\n你会怎么选？", COLORS['purple'], Inches(4.6), "📋"),
    ("函数越小越好？", "函数该写多长？\n拆得太小有问题吗？\n判断标准是什么？", COLORS['secondary'], Inches(8.4), "📏"),
]

for title, desc, color, x, icon in questions:
    y = Inches(2.0)
    add_rounded_rect(slide, x, y, Inches(3.7), Inches(3.5), COLORS['white'], color, 0.1)
    top_bar = add_shape(slide, MSO_SHAPE.RECTANGLE, x, y, Inches(3.7), Inches(0.7), color)
    add_text(slide, x + Inches(0.2), y + Inches(0.1), Inches(0.6), Inches(0.5), icon, font_size=24, align=PP_ALIGN.CENTER)
    add_text(slide, x + Inches(0.9), y + Inches(0.15), Inches(2.6), Inches(0.45), title, font_size=16, color=COLORS['white'], bold=True)
    add_text(slide, x + Inches(0.3), y + Inches(1.1), Inches(3.1), Inches(1.2),
             desc, font_size=13, color=COLORS['text'], align=PP_ALIGN.CENTER)
    add_rounded_rect(slide, x + Inches(0.3), y + Inches(2.6), Inches(3.1), Inches(0.6),
                     COLORS['light'], color, 0.3)
    add_text(slide, x + Inches(0.3), y + Inches(2.7), Inches(3.1), Inches(0.4),
             "💡 答案见课件文档", font_size=11, color=COLORS['text_light'], align=PP_ALIGN.CENTER)

encourage = add_rounded_rect(slide, Inches(1.5), Inches(5.9), Inches(10.3), Inches(1),
                             COLORS['yellow'], COLORS['yellow'], 0.3)
add_text(slide, Inches(1.5), Inches(6.0), Inches(10.3), Inches(0.4),
         "🌟 函数是封装的起点——后面所有的架构思想，都是在这个基础上一层一层堆起来的。",
         font_size=16, color=COLORS['dark'], bold=True, align=PP_ALIGN.CENTER)
add_text(slide, Inches(1.5), Inches(6.45), Inches(10.3), Inches(0.4),
         "带着问题学习下一讲——多文件编程，你会理解得更深刻！",
         font_size=13, color=COLORS['text_light'], align=PP_ALIGN.CENTER)

# ============================================================
# 第12页：小结
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "📝 小结", "函数是代码的家——封装、复用、隐藏细节")

gold_card = add_rounded_rect(slide, Inches(1.5), Inches(1.8), Inches(10.3), Inches(1.1),
                             COLORS['primary'], COLORS['primary'], 0.08)
add_text(slide, Inches(1.5), Inches(1.9), Inches(10.3), Inches(0.4),
         "🎯 一句话总结", font_size=15, color=COLORS['yellow'], bold=True, align=PP_ALIGN.CENTER)
add_text(slide, Inches(1.5), Inches(2.25), Inches(10.3), Inches(0.6),
         "函数是代码的家——把相关代码打包、命名、隐藏细节，让代码更清晰、更好维护、更容易复用。",
         font_size=16, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)

points = [
    ("1️⃣", "函数的意义", "封装代码\n消除重复\n提高复用", COLORS['blue']),
    ("2️⃣", "声明/定义/调用", "三阶段\n菜单→厨房→点菜", COLORS['accent']),
    ("3️⃣", "值传递", "传的是副本\n改副本不影响原件", COLORS['yellow']),
    ("4️⃣", "返回值", "return两个作用\n结束函数+返回结果", COLORS['purple']),
    ("5️⃣", "设计原则", "单一职责\n代码复用\n信息隐藏", COLORS['green']),
]

for i, (num, title, desc, color) in enumerate(points):
    x = Inches(0.5 + i * 2.55)
    y = Inches(3.4)
    add_rounded_rect(slide, x, y, Inches(2.35), Inches(2.5), COLORS['white'], color, 0.12)
    top_bar = add_shape(slide, MSO_SHAPE.RECTANGLE, x, y, Inches(2.35), Inches(0.08), color)
    add_text(slide, x, y + Inches(0.2), Inches(2.35), Inches(0.6), num, font_size=32, align=PP_ALIGN.CENTER)
    add_text(slide, x, y + Inches(0.9), Inches(2.35), Inches(0.45),
             title, font_size=17, color=COLORS['dark'], bold=True, align=PP_ALIGN.CENTER)
    divider = add_shape(slide, MSO_SHAPE.RECTANGLE, x + Inches(0.8), y + Inches(1.45), Inches(0.75), Inches(0.04), color)
    add_text(slide, x + Inches(0.15), y + Inches(1.6), Inches(2.05), Inches(0.8),
             desc, font_size=12, color=COLORS['text_light'], align=PP_ALIGN.CENTER)

next_card = add_rounded_rect(slide, Inches(3), Inches(6.2), Inches(7.3), Inches(0.8),
                             COLORS['accent'], COLORS['accent'], 0.3)
add_text(slide, Inches(3), Inches(6.3), Inches(7.3), Inches(0.6),
         "👉 下一讲：多文件编程 —— 函数太多了怎么办？拆成多个文件！",
         font_size=18, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)

# 保存
output_path = r"e:\2026-2027\2026-2027-1\AI化教学创新\AI化教学创新03\03_函数封装\docs\课件.pptx"
prs.save(output_path)
print(f"PPT生成完成：{output_path}")
print(f"共 {len(prs.slides)} 页")
