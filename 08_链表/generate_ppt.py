# -*- coding: utf-8 -*-
"""
第8讲：链表 —— PPT生成脚本
风格：轻松愉快、卡通风、明亮配色、有设计感
12页幻灯片，16:9，底部标注"第8讲 链表"
输出：docs/课件.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# ============================================================
# 配色方案（轻松活泼风，与系列一致）
# ============================================================
COLORS = {
    'primary':   RGBColor(0xFF, 0x8C, 0x42),    # 暖橙色
    'secondary':  RGBColor(0x4E, 0xCD, 0xC4),    # 薄荷绿
    'accent':     RGBColor(0xFF, 0x6B, 0x6B),    # 珊瑚红
    'dark':       RGBColor(0x2C, 0x3E, 0x50),    # 深灰蓝
    'light':      RGBColor(0xF7, 0xF9, 0xFC),    # 浅灰蓝
    'yellow':     RGBColor(0xFF, 0xD9, 0x3D),    # 明黄色
    'purple':     RGBColor(0xA2, 0x9B, 0xFE),    # 薰衣草紫
    'pink':       RGBColor(0xFF, 0x85, 0xA6),    # 粉红色
    'green':      RGBColor(0x6B, 0xCB, 0x77),    # 草绿色
    'blue':       RGBColor(0x74, 0xB9, 0xFF),    # 天蓝色
    'white':      RGBColor(0xFF, 0xFF, 0xFF),
    'text':       RGBColor(0x2C, 0x3E, 0x50),
    'text_light': RGBColor(0x7F, 0x8C, 0x8D),
    'code_bg':    RGBColor(0x1E, 0x29, 0x3B),    # 代码背景深色
    'code_header':RGBColor(0x34, 0x49, 0x5E),    # 代码标题栏
}

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# ============================================================
# 工具函数
# ============================================================

def add_bg(slide, color=COLORS['light']):
    """设置幻灯片背景"""
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg

def add_shape(slide, shape_type, left, top, width, height, fill_color=None, line_color=None, line_width=None):
    """通用形状创建"""
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    if fill_color is not None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if line_color is not None:
        shape.line.color.rgb = line_color
        if line_width:
            shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    return shape

def add_rounded_rect(slide, left, top, width, height, fill_color, border_color=None, radius=0.08):
    """圆角矩形"""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(2)
    else:
        shape.line.fill.background()
    shape.adjustments[0] = radius
    return shape

def add_circle(slide, left, top, size, fill_color):
    """圆形"""
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    return shape

def add_text(slide, left, top, width, height, text, font_size=24, color=COLORS['text'],
             bold=False, align=PP_ALIGN.LEFT, font_name='微软雅黑'):
    """文本框"""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = font_name
    return txBox

def add_code_block(slide, left, top, width, height, title, code_lines):
    """代码块（深色背景 + 标题栏 + 代码行）"""
    # 代码容器
    add_rounded_rect(slide, left, top, width, height, COLORS['code_bg'], COLORS['code_bg'], 0.06)
    # 标题栏
    add_shape(slide, MSO_SHAPE.RECTANGLE, left, top, width, Inches(0.4), COLORS['code_header'])
    # 小圆点
    for i, c in enumerate([COLORS['accent'], COLORS['yellow'], COLORS['green']]):
        add_circle(slide, left + Inches(0.15 + i * 0.3), top + Inches(0.08), Inches(0.18), c)
    add_text(slide, left + Inches(1.2), top + Inches(0.02), Inches(3), Inches(0.36),
             title, font_size=11, color=COLORS['text_light'])
    # 代码行
    for i, (code, color) in enumerate(code_lines):
        add_text(slide, left + Inches(0.4), top + Inches(0.5 + i * 0.32), width - Inches(0.6), Inches(0.3),
                 code, font_size=12, color=color, font_name='Consolas')

def add_title_bar(slide, title, subtitle=None):
    """标题栏（统一风格）"""
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLORS['primary']
    bar.line.fill.background()

    add_circle(slide, Inches(0.3), Inches(0.5), Inches(0.35), COLORS['primary'])
    add_circle(slide, Inches(0.55), Inches(0.55), Inches(0.25), COLORS['secondary'])

    add_text(slide, Inches(1.0), Inches(0.4), Inches(11), Inches(0.7),
             title, font_size=32, color=COLORS['dark'], bold=True)

    if subtitle:
        add_text(slide, Inches(1.0), Inches(1.05), Inches(11), Inches(0.45),
                 subtitle, font_size=15, color=COLORS['text_light'])

    # 底部标注
    add_text(slide, Inches(0.8), Inches(7.05), Inches(12), Inches(0.3),
             "第8讲 链表", font_size=11, color=COLORS['text_light'])

def draw_node(slide, left, top, width, height, data_text, fill_color, border_color=None):
    """画一个链表节点（数据域 + 指针域）"""
    # 整体背景
    node = add_rounded_rect(slide, left, top, width, height, COLORS['white'], border_color or fill_color, 0.08)
    # 数据域
    data_w = int(width * 0.65)
    add_shape(slide, MSO_SHAPE.RECTANGLE, left, top, data_w, height, fill_color)
    add_text(slide, left, top, data_w, height,
             data_text, font_size=13, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)
    # 指针域
    add_shape(slide, MSO_SHAPE.RECTANGLE, left + data_w, top, width - data_w, height,
              COLORS['light'], border_color or COLORS['text_light'])
    # 指针小圆点
    dot_size = Inches(0.15)
    add_circle(slide, left + data_w + (width - data_w) / 2 - dot_size / 2,
               top + height / 2 - dot_size / 2, dot_size, COLORS['accent'])
    # 标注 next
    add_text(slide, left + data_w, top + height - Inches(0.2), width - data_w, Inches(0.18),
             "next", font_size=8, color=COLORS['text_light'], align=PP_ALIGN.CENTER)

def draw_linked_list(slide, left, top, node_w, node_h, gap, items, colors):
    """画一整条链表（多个节点 + 箭头 + NULL结尾）"""
    data_w = int(node_w * 0.65)
    for i, (text, color) in enumerate(zip(items, colors)):
        x = left + i * (node_w + gap)
        draw_node(slide, x, top, node_w, node_h, text, color)
        # 箭头到下一个
        if i < len(items) - 1:
            arrow_x = x + node_w
            arrow_y = top + node_h / 2 - Inches(0.12)
            add_shape(slide, MSO_SHAPE.RIGHT_ARROW, arrow_x, arrow_y, gap, Inches(0.24),
                      COLORS['text_light'])
        else:
            # 最后一个 -> NULL
            null_x = x + node_w + Inches(0.1)
            add_text(slide, null_x, top + node_h / 2 - Inches(0.18), Inches(0.8), Inches(0.36),
                     "NULL", font_size=14, color=COLORS['text_light'], bold=True)
            # 斜杠
            add_shape(slide, MSO_SHAPE.RIGHT_ARROW, x + data_w, top + node_h / 2 - Inches(0.12),
                      node_w - data_w + Inches(0.1), Inches(0.24), COLORS['accent'])

def draw_train_car(slide, left, top, w, h, text, color):
    """画一节火车车厢"""
    # 车厢主体
    car = add_rounded_rect(slide, left, top, w, h, color, color, 0.1)
    # 窗户
    win_y = top + Inches(0.15)
    win_h = h * 0.35
    for i in range(2):
        add_rounded_rect(slide, left + Inches(0.15 + i * 0.65), win_y, Inches(0.5), win_h,
                         COLORS['white'], COLORS['white'], 0.2)
    # 文字
    add_text(slide, left, top + h * 0.55, w, h * 0.35,
             text, font_size=11, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)

def draw_train(slide, left, top, car_w, car_h, gap, items, colors):
    """画一列火车（车头 + 多节车厢 + 挂钩）"""
    # 车头
    head_w = car_w * 0.8
    add_rounded_rect(slide, left, top, head_w, car_h, COLORS['dark'], COLORS['dark'], 0.1)
    # 烟囱
    add_rounded_rect(slide, left + head_w * 0.2, top - Inches(0.25), Inches(0.2), Inches(0.3),
                     COLORS['dark'], COLORS['dark'], 0.2)
    add_circle(slide, left + head_w * 0.1, top - Inches(0.5), Inches(0.25), COLORS['text_light'])
    add_text(slide, left, top + car_h * 0.3, head_w, car_h * 0.4,
             "head", font_size=13, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)
    # 挂钩到第一节车厢
    add_shape(slide, MSO_SHAPE.RIGHT_ARROW, left + head_w, top + car_h / 2 - Inches(0.1),
              gap, Inches(0.2), COLORS['text_light'])

    x = left + head_w + gap
    for i, (text, color) in enumerate(zip(items, colors)):
        draw_train_car(slide, x, top, car_w, car_h, text, color)
        if i < len(items) - 1:
            add_shape(slide, MSO_SHAPE.RIGHT_ARROW, x + car_w, top + car_h / 2 - Inches(0.1),
                      gap, Inches(0.2), COLORS['text_light'])
        else:
            # 尾部 NULL
            add_text(slide, x + car_w + gap, top + car_h / 2 - Inches(0.18), Inches(0.8), Inches(0.36),
                     "NULL", font_size=13, color=COLORS['text_light'], bold=True)
        x += car_w + gap

# ============================================================
# 第1页：封面
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, COLORS['light'])

# 背景装饰圆
add_circle(slide, Inches(9.5), Inches(-1.5), Inches(5), COLORS['secondary'])
add_circle(slide, Inches(-2), Inches(4.5), Inches(4.5), COLORS['primary'])
add_circle(slide, Inches(11), Inches(5), Inches(2.5), COLORS['yellow'])

# 卡片
card = add_rounded_rect(slide, Inches(1.5), Inches(1.5), Inches(10.3), Inches(4.5),
                        COLORS['white'], COLORS['white'], 0.05)

# 左侧 - 火车比喻（链表比喻）
train_y = Inches(2.8)
train_colors = [COLORS['primary'], COLORS['secondary'], COLORS['yellow'], COLORS['green']]
draw_train(slide, Inches(2.0), train_y, Inches(0.7), Inches(0.9), Inches(0.25),
           ["A:100", "B:200", "C:500"], train_colors)
add_text(slide, Inches(1.8), Inches(4.2), Inches(4.5), Inches(0.4),
         "一列可以伸缩的火车 = 一个链表", font_size=14, color=COLORS['text_light'], align=PP_ALIGN.CENTER)

# 右侧文字
add_text(slide, Inches(5.8), Inches(2.3), Inches(6), Inches(1),
         "第8讲：链表", font_size=52, color=COLORS['dark'], bold=True)
add_text(slide, Inches(5.8), Inches(3.3), Inches(6), Inches(0.6),
         "动态数据的灵活管理", font_size=26, color=COLORS['primary'], bold=True)

# 分隔线
divider = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.8), Inches(4.1), Inches(1.5), Inches(0.06))
divider.fill.solid()
divider.fill.fore_color.rgb = COLORS['secondary']
divider.line.fill.background()

add_text(slide, Inches(5.8), Inches(4.3), Inches(6), Inches(0.5),
         "《从程序员到架构师》· C语言插件框架演进之旅", font_size=16, color=COLORS['text_light'])
add_text(slide, Inches(5.8), Inches(4.8), Inches(6), Inches(0.5),
         "单链表 · malloc/free · 增删改查 · 插件链表管理", font_size=18, color=COLORS['accent'], bold=True)

add_text(slide, Inches(6.0), Inches(5.4), Inches(2.3), Inches(0.5),
         "08 / 14", font_size=22, color=COLORS['accent'], bold=True, align=PP_ALIGN.CENTER)

add_text(slide, Inches(0.8), Inches(7.05), Inches(12), Inches(0.3),
         "第8讲 链表", font_size=11, color=COLORS['text_light'])

# ============================================================
# 第2页：知识图谱
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "知识图谱 · 第8讲", "本讲在整个知识体系中的位置")

# 中心节点
center_x = Inches(5.8)
center_y = Inches(3.7)

add_circle(slide, center_x - Inches(1.5), center_y - Inches(1), Inches(3), COLORS['secondary'])
add_circle(slide, center_x - Inches(1.1), center_y - Inches(0.6), Inches(2.2), COLORS['secondary'])

center_card = add_rounded_rect(slide, center_x - Inches(1.3), center_y - Inches(0.55), Inches(2.6), Inches(1.1),
                               COLORS['secondary'], COLORS['secondary'], 0.3)
add_text(slide, center_x - Inches(1.3), center_y - Inches(0.45), Inches(2.6), Inches(0.45),
         "链表", font_size=24, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)
add_text(slide, center_x - Inches(1.3), center_y + Inches(0.1), Inches(2.6), Inches(0.35),
         "（第8讲 核心）", font_size=13, color=COLORS['white'], align=PP_ALIGN.CENTER)

# 周围节点
nodes = [
    ("结构体", COLORS['purple'], Inches(1.0), Inches(1.3), "第7讲 · 前置", ""),
    ("指针", COLORS['yellow'], Inches(0.5), Inches(3.8), "第5讲 · 前置", ""),
    ("静态库", COLORS['green'], Inches(9.8), Inches(1.3), "→ 第9讲", ""),
    ("动态库", COLORS['blue'], Inches(9.8), Inches(3.8), "→ 第10讲", ""),
    ("插件链表管理", COLORS['accent'], Inches(5.3), Inches(6.0), "→ 第12讲 · 核心", ""),
]

for name, color, x, y, desc, icon in nodes:
    card = add_rounded_rect(slide, x, y, Inches(2.4), Inches(1.0),
                            COLORS['white'], color, 0.15)
    add_text(slide, x + Inches(0.1), y + Inches(0.15), Inches(0.6), Inches(0.7),
             icon, font_size=28, align=PP_ALIGN.CENTER)
    add_text(slide, x + Inches(0.7), y + Inches(0.1), Inches(1.6), Inches(0.4),
             name, font_size=15, color=COLORS['dark'], bold=True)
    add_text(slide, x + Inches(0.7), y + Inches(0.5), Inches(1.6), Inches(0.4),
             desc, font_size=10, color=COLORS['text_light'])

# 连接线
connections = [
    (center_x - Inches(0.5), center_y, Inches(3.0), Inches(1.8)),
    (center_x - Inches(0.5), center_y + Inches(0.2), Inches(2.0), Inches(4.3)),
    (center_x + Inches(0.5), center_y, Inches(9.8), Inches(1.8)),
    (center_x + Inches(0.5), center_y + Inches(0.2), Inches(9.8), Inches(4.3)),
    (center_x, center_y + Inches(0.5), Inches(6.5), Inches(6.5)),
]
for sx, sy, ex, ey in connections:
    line = slide.shapes.add_connector(1, sx, sy, ex, ey)
    line.line.color.rgb = COLORS['text_light']
    line.line.width = Pt(1.5)
    line.line.dash_style = 1

# 底部金句
tip_card = add_rounded_rect(slide, Inches(1.5), Inches(6.5), Inches(10.3), Inches(0.6),
                            COLORS['yellow'], COLORS['yellow'], 0.3)
add_text(slide, Inches(1.5), Inches(6.58), Inches(10.3), Inches(0.45),
         "链表是从'数据结构'到'软件架构'的关键桥梁——插件管理、配置化菜单都靠它",
         font_size=14, color=COLORS['dark'], bold=True, align=PP_ALIGN.CENTER)

# ============================================================
# 第3页：为什么需要链表（火车比喻）
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "为什么需要链表？", "从'固定储物柜'到'动态火车'")

# 左侧 - 固定数组的痛苦
left_x = Inches(0.6)
left_y = Inches(2.0)

pain_card = add_rounded_rect(slide, left_x, left_y, Inches(5.5), Inches(4.3),
                              COLORS['white'], COLORS['accent'], 0.08)
add_shape(slide, MSO_SHAPE.RECTANGLE, left_x, left_y, Inches(5.5), Inches(0.55), COLORS['accent'])
add_text(slide, left_x, left_y + Inches(0.08), Inches(5.5), Inches(0.4),
         "固定数组的痛苦", font_size=17, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)

# 固定格子
notes = [
    ("trans[0] = 100", COLORS['accent']),
    ("trans[1] = 200", COLORS['yellow']),
    ("trans[2] = 500", COLORS['secondary']),
    ("trans[3] = 300", COLORS['purple']),
    ("trans[4] = 150", COLORS['pink']),
    ("...最多10笔!", COLORS['accent']),
]
for i, (text, color) in enumerate(notes):
    ny = left_y + Inches(0.8 + i * 0.55)
    add_rounded_rect(slide, left_x + Inches(0.3), ny, Inches(4.9), Inches(0.45), color, color, 0.15)
    add_text(slide, left_x + Inches(0.3), ny + Inches(0.04), Inches(4.9), Inches(0.38),
             text, font_size=12, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)

add_text(slide, left_x + Inches(0.3), left_y + Inches(4.0), Inches(5.0), Inches(0.3),
         "想加第11笔？只能覆盖最旧的", font_size=13, color=COLORS['accent'], bold=True)

# 右侧 - 链表的自由
right_x = Inches(6.8)
right_y = Inches(2.0)

array_card = add_rounded_rect(slide, right_x, right_y, Inches(5.8), Inches(4.3),
                               COLORS['white'], COLORS['secondary'], 0.08)
add_shape(slide, MSO_SHAPE.RECTANGLE, right_x, right_y, Inches(5.8), Inches(0.55), COLORS['secondary'])
add_text(slide, right_x, right_y + Inches(0.08), Inches(5.8), Inches(0.4),
         "链表时代的自由", font_size=17, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)

# 火车图示
train_colors2 = [COLORS['accent'], COLORS['yellow'], COLORS['secondary'], COLORS['green']]
draw_train(slide, right_x + Inches(0.4), right_y + Inches(1.0), Inches(0.65), Inches(0.8), Inches(0.2),
           ["100", "200", "500", "..."], train_colors2)

add_text(slide, right_x + Inches(0.3), right_y + Inches(2.2), Inches(5.2), Inches(0.35),
         "head -> node -> node -> ... -> NULL",
         font_size=13, color=COLORS['secondary'], bold=True, align=PP_ALIGN.CENTER)
add_text(slide, right_x + Inches(0.3), right_y + Inches(2.7), Inches(5.2), Inches(0.35),
         "想加多少加多少，想删哪个删哪个", font_size=13, color=COLORS['text'])
add_text(slide, right_x + Inches(0.3), right_y + Inches(3.2), Inches(5.2), Inches(0.35),
         "malloc 按需分配，不受固定大小限制", font_size=13, color=COLORS['text'])
add_text(slide, right_x + Inches(0.3), right_y + Inches(3.7), Inches(5.2), Inches(0.35),
         "指针串联节点，大小随时变化", font_size=13, color=COLORS['text'])

# 中间箭头
arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(6.0), Inches(3.8), Inches(0.8), Inches(0.5))
arrow.fill.solid()
arrow.fill.fore_color.rgb = COLORS['primary']
arrow.line.fill.background()

# 底部金句
gold = add_rounded_rect(slide, Inches(1.5), Inches(6.5), Inches(10.3), Inches(0.6),
                        COLORS['primary'], COLORS['primary'], 0.3)
add_text(slide, Inches(1.5), Inches(6.58), Inches(10.3), Inches(0.45),
         "金句：数组是固定的储物柜，链表是可以伸缩的火车——大小随需变化",
         font_size=15, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)

# ============================================================
# 第4页：单链表结构
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "单链表结构", "数据域 + 指针域 = 火车车厢")

# 上方 - 节点结构图
add_text(slide, Inches(0.8), Inches(1.7), Inches(5), Inches(0.4),
         "一个链表节点 = 数据域 + 指针域", font_size=16, color=COLORS['dark'], bold=True)

# 画一个放大节点
node_x = Inches(2.0)
node_y = Inches(2.3)
node_w = Inches(5.5)
node_h = Inches(1.2)

# 数据域
data_card = add_rounded_rect(slide, node_x, node_y, Inches(3.5), node_h,
                              COLORS['primary'], COLORS['primary'], 0.06)
add_text(slide, node_x, node_y + Inches(0.1), Inches(3.5), Inches(0.4),
         "数据域", font_size=14, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)
add_text(slide, node_x, node_y + Inches(0.55), Inches(3.5), Inches(0.35),
         "double amount", font_size=12, color=COLORS['white'], align=PP_ALIGN.CENTER)
add_text(slide, node_x, node_y + Inches(0.85), Inches(3.5), Inches(0.35),
         "int type / char desc[]", font_size=12, color=COLORS['white'], align=PP_ALIGN.CENTER)

# 指针域
ptr_card = add_rounded_rect(slide, node_x + Inches(3.5), node_y, Inches(2.0), node_h,
                             COLORS['secondary'], COLORS['secondary'], 0.06)
add_text(slide, node_x + Inches(3.5), node_y + Inches(0.1), Inches(2.0), Inches(0.4),
         "指针域", font_size=14, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)
add_text(slide, node_x + Inches(3.5), node_y + Inches(0.55), Inches(2.0), Inches(0.35),
         "next", font_size=12, color=COLORS['white'], align=PP_ALIGN.CENTER)
add_text(slide, node_x + Inches(3.5), node_y + Inches(0.85), Inches(2.0), Inches(0.35),
         "-> 下一个", font_size=12, color=COLORS['white'], align=PP_ALIGN.CENTER)

# 右侧代码
code_lines = [
    ('typedef struct Node {', COLORS['blue']),
    ('  /* 数据域 */', COLORS['text_light']),
    ('  double amount;', COLORS['green']),
    ('  int    type;', COLORS['green']),
    ('  /* 指针域 */', COLORS['text_light']),
    ('  struct Node* next;', COLORS['yellow']),
    ('} Node;', COLORS['blue']),
]
add_code_block(slide, Inches(8.5), Inches(2.0), Inches(4.2), Inches(3.2),
               "linked_list.c", code_lines)

# 下方 - 完整链表图示
add_text(slide, Inches(0.8), Inches(4.0), Inches(5), Inches(0.4),
         "完整链表：火车车厢串联", font_size=16, color=COLORS['dark'], bold=True)

# head 指针标注
add_text(slide, Inches(0.5), Inches(4.8), Inches(0.8), Inches(0.4),
         "head", font_size=14, color=COLORS['accent'], bold=True)
add_shape(slide, MSO_SHAPE.RIGHT_ARROW, Inches(1.1), Inches(4.9), Inches(0.4), Inches(0.2), COLORS['accent'])

# 画3个节点
ll_colors = [COLORS['primary'], COLORS['secondary'], COLORS['green']]
draw_linked_list(slide, Inches(1.6), Inches(4.5), Inches(1.8), Inches(0.8), Inches(0.5),
                 ["A:100", "B:200", "C:500"], ll_colors)

# 底部金句
gold = add_rounded_rect(slide, Inches(1.5), Inches(6.3), Inches(10.3), Inches(0.6),
                        COLORS['secondary'], COLORS['secondary'], 0.3)
add_text(slide, Inches(1.5), Inches(6.38), Inches(10.3), Inches(0.45),
         "比喻：数据域 = 车厢里的货物，指针域 = 车厢之间的挂钩",
         font_size=15, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)

# ============================================================
# 第5页：创建与销毁（malloc/free）
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "创建与销毁", "malloc 分配内存 / free 释放内存")

# 左侧 - 创建节点
left_x = Inches(0.6)
create_card = add_rounded_rect(slide, left_x, Inches(1.8), Inches(6.0), Inches(4.5),
                               COLORS['white'], COLORS['green'], 0.08)
add_shape(slide, MSO_SHAPE.RECTANGLE, left_x, Inches(1.8), Inches(6.0), Inches(0.5), COLORS['green'])
add_text(slide, left_x, Inches(1.85), Inches(6.0), Inches(0.4),
         "创建节点：malloc", font_size=16, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)

create_code = [
    ('Node* create(double amt) {', COLORS['blue']),
    ('  // 1. malloc 分配内存', COLORS['text_light']),
    ('  Node* n = malloc(sizeof(Node));', COLORS['green']),
    ('', COLORS['text']),
    ('  // 2. 检查是否成功', COLORS['text_light']),
    ('  if (n == NULL) return NULL;', COLORS['yellow']),
    ('', COLORS['text']),
    ('  // 3. 填充数据域', COLORS['text_light']),
    ('  n->amount = amt;', COLORS['green']),
    ('  n->next = NULL; // 4. next置空', COLORS['green']),
    ('  return n;', COLORS['blue']),
    ('}', COLORS['blue']),
]
add_code_block(slide, left_x + Inches(0.2), Inches(2.5), Inches(5.6), Inches(3.6),
               "create_node.c", create_code)

# 右侧 - 销毁链表
right_x = Inches(7.0)
destroy_card = add_rounded_rect(slide, right_x, Inches(1.8), Inches(5.8), Inches(4.5),
                                COLORS['white'], COLORS['accent'], 0.08)
add_shape(slide, MSO_SHAPE.RECTANGLE, right_x, Inches(1.8), Inches(5.8), Inches(0.5), COLORS['accent'])
add_text(slide, right_x, Inches(1.85), Inches(5.8), Inches(0.4),
         "销毁链表：free（防内存泄漏）", font_size=16, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)

destroy_code = [
    ('void destroy(Node* head) {', COLORS['blue']),
    ('  Node* cur = head;', COLORS['green']),
    ('  while (cur != NULL) {', COLORS['blue']),
    ('    // 先保存下一个！', COLORS['accent']),
    ('    Node* tmp = cur;', COLORS['yellow']),
    ('    cur = cur->next;', COLORS['yellow']),
    ('    free(tmp); // 再释放', COLORS['green']),
    ('  }', COLORS['blue']),
    ('}', COLORS['blue']),
    ('', COLORS['text']),
    ('// 错误！先free再访问next:', COLORS['accent']),
    ('// free(cur); cur=cur->next;', COLORS['accent']),
    ('// 未定义行为！', COLORS['accent']),
]
add_code_block(slide, right_x + Inches(0.2), Inches(2.5), Inches(5.4), Inches(3.8),
               "destroy_list.c", destroy_code)

# 底部金句
gold = add_rounded_rect(slide, Inches(1.5), Inches(6.5), Inches(10.3), Inches(0.6),
                        COLORS['yellow'], COLORS['yellow'], 0.3)
add_text(slide, Inches(1.5), Inches(6.58), Inches(10.3), Inches(0.45),
         "铁律：每个 malloc 都要有对应的 free！先保存 next 再 free 当前节点",
         font_size=14, color=COLORS['dark'], bold=True, align=PP_ALIGN.CENTER)

# ============================================================
# 第6页：头部插入
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "增 · 头部插入", "新节点插到最前面——O(1) 高效操作")

# 上方 - 图解
add_text(slide, Inches(0.8), Inches(1.6), Inches(5), Inches(0.4),
         "插入前：", font_size=14, color=COLORS['text'], bold=True)

# 原链表
ll_colors2 = [COLORS['primary'], COLORS['secondary'], COLORS['green']]
draw_linked_list(slide, Inches(1.5), Inches(2.1), Inches(1.5), Inches(0.7), Inches(0.5),
                 ["A", "B", "C"], ll_colors2)

add_text(slide, Inches(0.8), Inches(3.1), Inches(5), Inches(0.4),
         "插入后（NEW 插到头部）：", font_size=14, color=COLORS['accent'], bold=True)

# 新链表
new_colors = [COLORS['accent']] + ll_colors2
draw_linked_list(slide, Inches(0.8), Inches(3.6), Inches(1.5), Inches(0.7), Inches(0.4),
                 ["NEW", "A", "B", "C"], new_colors)

# 右侧代码
code_lines = [
    ('void insert_head(double amt) {', COLORS['blue']),
    ('  // 1. 创建新节点', COLORS['text_light']),
    ('  Node* n = create(amt);', COLORS['green']),
    ('', COLORS['text']),
    ('  // 2. 新节点指向旧头', COLORS['text_light']),
    ('  n->next = head;', COLORS['yellow']),
    ('', COLORS['text']),
    ('  // 3. 头指针指向新节点', COLORS['text_light']),
    ('  head = n;', COLORS['yellow']),
    ('}', COLORS['blue']),
    ('', COLORS['text']),
    ('// 时间复杂度: O(1)!', COLORS['green']),
    ('// 只改2个指针，无需遍历', COLORS['green']),
]
add_code_block(slide, Inches(7.5), Inches(1.8), Inches(5.3), Inches(4.0),
               "insert_head.c", code_lines)

# 步骤标注
steps = [
    ("1.创建", COLORS['green']),
    ("2.n->next=head", COLORS['yellow']),
    ("3.head=n", COLORS['accent']),
]
for i, (text, color) in enumerate(steps):
    x = Inches(1.0 + i * 2.0)
    add_rounded_rect(slide, x, Inches(4.6), Inches(1.8), Inches(0.4), color, color, 0.2)
    add_text(slide, x, Inches(4.63), Inches(1.8), Inches(0.35),
             text, font_size=11, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)

# 底部金句
gold = add_rounded_rect(slide, Inches(1.5), Inches(5.5), Inches(10.3), Inches(0.6),
                        COLORS['primary'], COLORS['primary'], 0.3)
add_text(slide, Inches(1.5), Inches(5.58), Inches(10.3), Inches(0.45),
         "头部插入 O(1) vs 数组头部插入 O(n)——链表完胜！",
         font_size=15, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)

# ============================================================
# 第7页：尾部插入
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "增 · 尾部插入", "新节点追加到末尾——需要遍历 O(n)")

# 上方 - 图解
add_text(slide, Inches(0.8), Inches(1.6), Inches(5), Inches(0.4),
         "插入前：", font_size=14, color=COLORS['text'], bold=True)

ll_colors3 = [COLORS['primary'], COLORS['secondary'], COLORS['green']]
draw_linked_list(slide, Inches(1.5), Inches(2.1), Inches(1.5), Inches(0.7), Inches(0.5),
                 ["A", "B", "C"], ll_colors3)

add_text(slide, Inches(0.8), Inches(3.1), Inches(5), Inches(0.4),
         "插入后（NEW 追加到尾部）：", font_size=14, color=COLORS['accent'], bold=True)

new_colors2 = ll_colors3 + [COLORS['accent']]
draw_linked_list(slide, Inches(0.8), Inches(3.6), Inches(1.5), Inches(0.7), Inches(0.4),
                 ["A", "B", "C", "NEW"], new_colors2)

# 右侧代码
code_lines = [
    ('void insert_tail(double amt) {', COLORS['blue']),
    ('  Node* n = create(amt);', COLORS['green']),
    ('', COLORS['text']),
    ('  // 情况1: 链表为空', COLORS['text_light']),
    ('  if (head == NULL) {', COLORS['yellow']),
    ('    head = n;', COLORS['yellow']),
    ('    return;', COLORS['yellow']),
    ('  }', COLORS['blue']),
    ('', COLORS['text']),
    ('  // 情况2: 遍历到最后', COLORS['text_light']),
    ('  Node* cur = head;', COLORS['green']),
    ('  while (cur->next != NULL)', COLORS['blue']),
    ('    cur = cur->next;', COLORS['blue']),
    ('  cur->next = n; // 挂上新节点', COLORS['green']),
    ('}', COLORS['blue']),
    ('// 时间复杂度: O(n)', COLORS['accent']),
]
add_code_block(slide, Inches(7.5), Inches(1.8), Inches(5.3), Inches(4.5),
               "insert_tail.c", code_lines)

# 底部金句
gold = add_rounded_rect(slide, Inches(1.5), Inches(6.0), Inches(10.3), Inches(0.6),
                        COLORS['yellow'], COLORS['yellow'], 0.3)
add_text(slide, Inches(1.5), Inches(6.08), Inches(10.3), Inches(0.45),
         "优化：维护尾指针 tail 可让尾部插入也达到 O(1)",
         font_size=14, color=COLORS['dark'], bold=True, align=PP_ALIGN.CENTER)

# ============================================================
# 第8页：遍历与查找
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "查 · 遍历与查找", "从头到 NULL 逐个访问——O(n)")

# 上方 - 遍历图示
add_text(slide, Inches(0.8), Inches(1.6), Inches(5), Inches(0.4),
         "遍历：cur 从 head 移动到 NULL", font_size=14, color=COLORS['text'], bold=True)

# 画链表 + cur 指针移动
ll_colors4 = [COLORS['primary'], COLORS['secondary'], COLORS['green'], COLORS['purple']]
draw_linked_list(slide, Inches(1.0), Inches(2.1), Inches(1.4), Inches(0.65), Inches(0.4),
                 ["A", "B", "C", "D"], ll_colors4)

# cur 移动标注
add_text(slide, Inches(1.2), Inches(2.9), Inches(1.4), Inches(0.3),
         "cur", font_size=11, color=COLORS['accent'], bold=True, align=PP_ALIGN.CENTER)
add_shape(slide, MSO_SHAPE.DOWN_ARROW, Inches(1.7), Inches(2.85), Inches(0.3), Inches(0.2), COLORS['accent'])

# 右侧 - 代码对比
code_lines = [
    ('// 第6讲: 数组遍历', COLORS['text_light']),
    ('for (i=0; i<10; i++)', COLORS['blue']),
    ('  printf("%f", arr[i]);', COLORS['green']),
    ('', COLORS['text']),
    ('// 第8讲: 链表遍历', COLORS['text_light']),
    ('Node* cur = head;', COLORS['green']),
    ('while (cur != NULL) {', COLORS['blue']),
    ('  printf("%.2f", cur->amount);', COLORS['green']),
    ('  cur = cur->next; // 前进!', COLORS['yellow']),
    ('}', COLORS['blue']),
    ('', COLORS['text']),
    ('// 查找: 遍历中比较', COLORS['text_light']),
    ('Node* find(int id) {', COLORS['blue']),
    ('  while(cur) {', COLORS['blue']),
    ('    if(cur->id == id) return cur;', COLORS['green']),
    ('    cur = cur->next;', COLORS['yellow']),
    ('  }', COLORS['blue']),
    ('  return NULL; // 没找到', COLORS['accent']),
    ('}', COLORS['blue']),
]
add_code_block(slide, Inches(7.0), Inches(1.8), Inches(5.8), Inches(5.0),
               "traverse_find.c", code_lines)

# 底部金句
gold = add_rounded_rect(slide, Inches(1.0), Inches(6.4), Inches(6.0), Inches(0.5),
                        COLORS['secondary'], COLORS['secondary'], 0.3)
add_text(slide, Inches(1.0), Inches(6.47), Inches(6.0), Inches(0.35),
         "数组: for + 索引 i    链表: while + 指针 cur->next",
         font_size=13, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)

# ============================================================
# 第9页：删除操作
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "删 · 按值删除", "找到前驱，跳过目标，free 回收")

# 上方 - 图解
add_text(slide, Inches(0.8), Inches(1.6), Inches(5), Inches(0.4),
         "删除前（要删 B）：", font_size=14, color=COLORS['text'], bold=True)

del_colors = [COLORS['primary'], COLORS['accent'], COLORS['green']]
draw_linked_list(slide, Inches(1.0), Inches(2.1), Inches(1.4), Inches(0.65), Inches(0.4),
                 ["A", "B(删)", "C"], del_colors)

add_text(slide, Inches(0.8), Inches(3.1), Inches(5), Inches(0.4),
         "删除后（B 被 free 回收）：", font_size=14, color=COLORS['accent'], bold=True)

after_colors = [COLORS['primary'], COLORS['green']]
draw_linked_list(slide, Inches(1.0), Inches(3.6), Inches(1.4), Inches(0.65), Inches(0.4),
                 ["A", "C"], after_colors)

# 关键：前驱标注
add_text(slide, Inches(1.0), Inches(4.5), Inches(1.4), Inches(0.3),
         "prev", font_size=11, color=COLORS['yellow'], bold=True, align=PP_ALIGN.CENTER)
add_shape(slide, MSO_SHAPE.DOWN_ARROW, Inches(1.5), Inches(4.45), Inches(0.3), Inches(0.2), COLORS['yellow'])

# 右侧代码
code_lines = [
    ('void delete(double amt) {', COLORS['blue']),
    ('  if (head == NULL) return;', COLORS['text_light']),
    ('', COLORS['text']),
    ('  // 情况1: 删头节点', COLORS['text_light']),
    ('  if (head->amount == amt) {', COLORS['yellow']),
    ('    Node* tmp = head;', COLORS['green']),
    ('    head = head->next;', COLORS['green']),
    ('    free(tmp); return;', COLORS['accent']),
    ('  }', COLORS['blue']),
    ('', COLORS['text']),
    ('  // 情况2: 删中间/末尾', COLORS['text_light']),
    ('  Node* prev = head;', COLORS['green']),
    ('  Node* cur = head->next;', COLORS['green']),
    ('  while (cur != NULL) {', COLORS['blue']),
    ('    if (cur->amount == amt) {', COLORS['yellow']),
    ('      prev->next = cur->next;', COLORS['yellow']),
    ('      free(cur); return;', COLORS['accent']),
    ('    }', COLORS['blue']),
    ('    prev = cur;', COLORS['green']),
    ('    cur = cur->next;', COLORS['green']),
    ('  }', COLORS['blue']),
    ('}', COLORS['blue']),
]
add_code_block(slide, Inches(6.8), Inches(1.8), Inches(6.0), Inches(5.0),
               "delete_node.c", code_lines)

# 底部金句
gold = add_rounded_rect(slide, Inches(1.5), Inches(6.4), Inches(10.3), Inches(0.5),
                        COLORS['accent'], COLORS['accent'], 0.3)
add_text(slide, Inches(1.5), Inches(6.47), Inches(10.3), Inches(0.35),
         "关键：删除需要保存前驱节点！prev->next = cur->next 跳过被删节点",
         font_size=14, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)

# ============================================================
# 第10页：链表 vs 数组对比
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "链表 vs 数组", "各有优劣，根据场景选择")

# 左侧 - 数组
left_x = Inches(0.6)
arr_card = add_rounded_rect(slide, left_x, Inches(1.8), Inches(5.8), Inches(4.5),
                             COLORS['white'], COLORS['blue'], 0.08)
add_shape(slide, MSO_SHAPE.RECTANGLE, left_x, Inches(1.8), Inches(5.8), Inches(0.5), COLORS['blue'])
add_text(slide, left_x, Inches(1.85), Inches(5.8), Inches(0.4),
         "数组", font_size=18, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)

arr_props = [
    ("大小", "固定，编译时确定", COLORS['green']),
    ("内存", "连续分配", COLORS['green']),
    ("随机访问", "arr[i]  O(1)", COLORS['green']),
    ("头部插入", "全部后移  O(n)", COLORS['accent']),
    ("删除", "全部前移  O(n)", COLORS['accent']),
    ("额外开销", "无", COLORS['green']),
    ("缓存友好", "是（连续内存）", COLORS['green']),
]
for i, (prop, val, color) in enumerate(arr_props):
    y = Inches(2.5 + i * 0.5)
    add_text(slide, left_x + Inches(0.3), y, Inches(1.8), Inches(0.35),
             prop, font_size=13, color=COLORS['dark'], bold=True)
    add_text(slide, left_x + Inches(2.2), y, Inches(3.3), Inches(0.35),
             val, font_size=13, color=color)

# 右侧 - 链表
right_x = Inches(6.9)
list_card = add_rounded_rect(slide, right_x, Inches(1.8), Inches(5.8), Inches(4.5),
                              COLORS['white'], COLORS['secondary'], 0.08)
add_shape(slide, MSO_SHAPE.RECTANGLE, right_x, Inches(1.8), Inches(5.8), Inches(0.5), COLORS['secondary'])
add_text(slide, right_x, Inches(1.85), Inches(5.8), Inches(0.4),
         "链表", font_size=18, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)

list_props = [
    ("大小", "动态，随时变化", COLORS['green']),
    ("内存", "分散分配，指针串联", COLORS['yellow']),
    ("随机访问", "必须遍历  O(n)", COLORS['accent']),
    ("头部插入", "改2指针  O(1)", COLORS['green']),
    ("删除", "改指针  O(1)", COLORS['green']),
    ("额外开销", "每节点多1个next", COLORS['accent']),
    ("缓存友好", "否（分散内存）", COLORS['accent']),
]
for i, (prop, val, color) in enumerate(list_props):
    y = Inches(2.5 + i * 0.5)
    add_text(slide, right_x + Inches(0.3), y, Inches(1.8), Inches(0.35),
             prop, font_size=13, color=COLORS['dark'], bold=True)
    add_text(slide, right_x + Inches(2.2), y, Inches(3.3), Inches(0.35),
             val, font_size=13, color=color)

# 底部金句
gold = add_rounded_rect(slide, Inches(1.5), Inches(6.5), Inches(10.3), Inches(0.6),
                        COLORS['primary'], COLORS['primary'], 0.3)
add_text(slide, Inches(1.5), Inches(6.58), Inches(10.3), Inches(0.45),
         "选择原则：大小已知+频繁随机访问 → 数组 | 大小动态+频繁增删 → 链表",
         font_size=14, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)

# ============================================================
# 第11页：插件链表管理（预告）
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "插件链表管理（预告第12-13讲）", "链表是从数据结构到架构的桥梁")

# 上方 - 插件链表示意图
add_text(slide, Inches(0.8), Inches(1.6), Inches(6), Inches(0.4),
         "每加载一个插件 = 链表加一个节点", font_size=15, color=COLORS['dark'], bold=True)

# 插件链表
plugin_colors = [COLORS['purple'], COLORS['blue'], COLORS['secondary']]
plugin_names = ["汇率查询", "积分兑换", "电子发票"]
draw_linked_list(slide, Inches(1.0), Inches(2.2), Inches(2.0), Inches(0.8), Inches(0.5),
                 plugin_names, plugin_colors)

# 加载标注
add_text(slide, Inches(1.0), Inches(3.2), Inches(7), Inches(0.35),
         "load_plugin() → 头部插入到插件链表", font_size=12, color=COLORS['purple'], bold=True)
add_text(slide, Inches(1.0), Inches(3.6), Inches(7), Inches(0.35),
         "run_all_plugins() → 遍历链表依次执行", font_size=12, color=COLORS['green'], bold=True)

# 右侧 - 配置化菜单预告
right_x = Inches(8.5)
menu_card = add_rounded_rect(slide, right_x, Inches(1.8), Inches(4.3), Inches(4.2),
                              COLORS['white'], COLORS['accent'], 0.08)
add_shape(slide, MSO_SHAPE.RECTANGLE, right_x, Inches(1.8), Inches(4.3), Inches(0.5), COLORS['accent'])
add_text(slide, right_x, Inches(1.85), Inches(4.3), Inches(0.4),
         "配置化菜单（第13讲）", font_size=15, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)

menu_items = [
    "1. 查询余额  (内置)",
    "2. 存款      (内置)",
    "3. 取款      (内置)",
    "4. 汇率查询  (插件加载)",
    "5. 积分兑换  (插件加载)",
    "6. 电子发票  (插件加载)",
]
for i, item in enumerate(menu_items):
    color = COLORS['green'] if "内置" in item else COLORS['purple']
    add_text(slide, right_x + Inches(0.3), Inches(2.5 + i * 0.45), Inches(3.8), Inches(0.35),
             item, font_size=12, color=color, bold=True)

add_text(slide, right_x + Inches(0.3), Inches(5.3), Inches(3.8), Inches(0.35),
         "菜单项 = 链表节点", font_size=11, color=COLORS['accent'], bold=True)
add_text(slide, right_x + Inches(0.3), Inches(5.65), Inches(3.8), Inches(0.35),
         "加载插件自动加菜单项!", font_size=11, color=COLORS['accent'], bold=True)

# 下方 - 代码
code_lines = [
    ('typedef struct PluginNode {', COLORS['blue']),
    ('  char name[20];           // 插件名', COLORS['green']),
    ('  void (*handler)();       // 函数指针!', COLORS['yellow']),
    ('  struct PluginNode* next; // 链表串联', COLORS['secondary']),
    ('} PluginNode;', COLORS['blue']),
]
add_code_block(slide, Inches(0.8), Inches(4.5), Inches(7.0), Inches(2.2),
               "plugin.h (第12讲预告)", code_lines)

# 底部金句
gold = add_rounded_rect(slide, Inches(1.5), Inches(6.8), Inches(10.3), Inches(0.5),
                        COLORS['purple'], COLORS['purple'], 0.3)
add_text(slide, Inches(1.5), Inches(6.87), Inches(10.3), Inches(0.35),
         "链表 + 函数指针 = 插件框架的核心机制！",
         font_size=14, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)

# ============================================================
# 第12页：小结
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_title_bar(slide, "小结", "从固定数组到动态链表")

# 五大收获
gains = [
    ("动态大小", "想加多少加多少\n不受固定上限", COLORS['green']),
    ("随时增删", "运行时动态增删\nmalloc/free 管理", COLORS['secondary']),
    ("指针实战", "next 指针串联节点\n指针是链表的灵魂", COLORS['primary']),
    ("插件基石", "插件链表管理\n配置化菜单的基础", COLORS['purple']),
    ("架构桥梁", "从数据结构\n到软件架构的关键一步", COLORS['accent']),
]

for i, (title, desc, color) in enumerate(gains):
    x = Inches(0.5 + i * 2.55)
    y = Inches(1.8)
    card = add_rounded_rect(slide, x, y, Inches(2.3), Inches(2.5),
                            COLORS['white'], color, 0.1)
    add_shape(slide, MSO_SHAPE.RECTANGLE, x, y, Inches(2.3), Inches(0.08), color)
    add_text(slide, x + Inches(0.1), y + Inches(0.2), Inches(2.1), Inches(0.5),
             title, font_size=15, color=color, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, x + Inches(0.15), y + Inches(0.8), Inches(2.0), Inches(1.5),
             desc, font_size=12, color=COLORS['text'], align=PP_ALIGN.CENTER)

# 中间 - ATM演进对比
add_text(slide, Inches(0.8), Inches(4.6), Inches(6), Inches(0.4),
         "ATM 案例演进：", font_size=15, color=COLORS['dark'], bold=True)

evolution = [
    ("第6讲", "trans[10] 固定数组", COLORS['blue']),
    ("第7讲", "结构体数组", COLORS['purple']),
    ("第8讲", "链表动态无限!", COLORS['green']),
]
for i, (lecture, desc, color) in enumerate(evolution):
    x = Inches(0.8 + i * 4.0)
    y = Inches(5.1)
    add_rounded_rect(slide, x, y, Inches(3.5), Inches(0.7), color, color, 0.15)
    add_text(slide, x + Inches(0.1), y + Inches(0.05), Inches(3.3), Inches(0.3),
             lecture, font_size=12, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, x + Inches(0.1), y + Inches(0.35), Inches(3.3), Inches(0.3),
             desc, font_size=11, color=COLORS['white'], align=PP_ALIGN.CENTER)
    if i < 2:
        add_shape(slide, MSO_SHAPE.RIGHT_ARROW, x + Inches(3.5), y + Inches(0.2),
                  Inches(0.4), Inches(0.3), COLORS['text_light'])

# 底部金句
gold = add_rounded_rect(slide, Inches(1.5), Inches(6.2), Inches(10.3), Inches(0.8),
                        COLORS['primary'], COLORS['primary'], 0.2)
add_text(slide, Inches(1.5), Inches(6.3), Inches(10.3), Inches(0.35),
         "链表不仅仅是一种数据结构，更是从'写代码'到'做架构'的关键桥梁",
         font_size=15, color=COLORS['white'], bold=True, align=PP_ALIGN.CENTER)
add_text(slide, Inches(1.5), Inches(6.65), Inches(10.3), Inches(0.3),
         "第12讲：插件链表管理 | 第13讲：配置化菜单",
         font_size=13, color=COLORS['white'], align=PP_ALIGN.CENTER)

# ============================================================
# 保存
# ============================================================
output_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'docs')
os.makedirs(output_dir, exist_ok=True)
output_path = os.path.join(output_dir, '课件.pptx')
prs.save(output_path)
print(f"PPT已生成：{output_path}")
print(f"共 12 页幻灯片")
